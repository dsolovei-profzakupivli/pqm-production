"""Real HTTP export against a disposable on-disk copy, not connection= shortcuts."""
from contextlib import closing
from datetime import datetime
import gc
import hashlib
import http.client
from http.server import HTTPServer
import io
from pathlib import Path
import sqlite3
import tempfile
import threading
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
import server
import supplier_xlsx
import test_prod_framework_registry as fixtures


class SupplierExcelEndpointTests(fixtures.SupplierFilterExcelTests):
    def setUp(self):
        super().setUp()
        self.add('first', code='0001', name='=literal\r\nУкраїна')
        self.add('second', code='0002', day=2, contract='terminated')
        self.con.commit()
        self.folder = tempfile.TemporaryDirectory(prefix='pqm-export-http-')
        self.addCleanup(self.folder.cleanup)
        self.path = Path(self.folder.name) / 'isolated.sqlite3'
        with closing(sqlite3.connect(self.path)) as copied:
            self.con.backup(copied)
        for key, value in (('DB_PATH', self.path), ('AUTH_ENABLED', False),
                           ('IS_WEB_ENV', False), ('SANDBOX_MODE', False)):
            context = patch.object(server, key, value)
            context.start()
            self.addCleanup(context.stop)
        self.http = HTTPServer(('127.0.0.1', 0), server.Handler)
        self.thread = threading.Thread(target=self.http.serve_forever, daemon=True)
        self.thread.start()
        self.addCleanup(self.stop_http)

    def stop_http(self):
        self.http.shutdown()
        self.http.server_close()
        self.thread.join(timeout=5)
        gc.collect()  # Existing db() transaction contexts do not close connections.

    def get(self, query):
        with closing(http.client.HTTPConnection('127.0.0.1', self.http.server_port, timeout=10)) as client:
            client.request('GET', '/api/suppliers-registry/export.xlsx?' + query)
            response = client.getresponse()
            return response.status, dict(response.getheaders()), response.read()

    def snapshot(self):
        with closing(sqlite3.connect(self.path.as_uri() + '?mode=ro', uri=True)) as con:
            return hashlib.sha256('\n'.join(con.iterdump()).encode()).hexdigest()

    def test_real_endpoint_opens_database_and_returns_v5_workbook(self):
        before = self.snapshot()
        # Do not mock either export function or supply an existing connection.
        with patch.object(server, 'supplier_registry_excel', wraps=server.supplier_registry_excel) as export:
            status, headers, raw = self.get('registry_only=1&framework_id=f1')
            export.assert_called_once_with({'registry_only': ['1'], 'framework_id': ['f1']})
        self.assertEqual(status, 200)
        self.assertEqual(headers['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertEqual(headers['Content-Disposition'], 'attachment; filename="03110000-1_pqm-suppliers.xlsx"')
        self.assertEqual(int(headers['Content-Length']), len(raw))
        self.assertEqual(headers['Cache-Control'], 'no-store')
        self.assertTrue(raw.startswith(b'PK'))
        ws = load_workbook(io.BytesIO(raw)).active
        self.assertEqual([cell.value for cell in ws[1]], list(supplier_xlsx.HEADERS))
        self.assertEqual((ws.max_row, ws.max_column), (3, 6))
        self.assertEqual([ws.cell(row, 1).value for row in (2, 3)], ['0002', '0001'])
        self.assertEqual(ws['B3'].value, '=literal\r\nУкраїна')
        self.assertEqual(ws['B3'].data_type, 's')
        self.assertEqual(ws['C2'].value, datetime(2026, 10, 2))
        self.assertEqual(ws['C2'].number_format, 'dd.mm.yyyy')
        self.assertEqual(ws['D2'].value, 'Неактивний')
        self.assertEqual(ws['E3'].value, 'first@example.invalid')
        self.assertEqual(ws['F3'].value, '+380001234567')
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertEqual(ws.tables['PqmSuppliers'].autoFilter.ref, 'A1:F3')
        self.assertEqual(ws['A1'].fill.fgColor.rgb, 'FFDCE6F1')
        self.assertEqual(ws['A2'].font.color.rgb, 'FF000000')
        self.assertEqual(ws['A2'].border.bottom.color.rgb, 'FFD0DAE3')
        ns = {'s': supplier_xlsx.NS}
        with ZipFile(io.BytesIO(raw)) as archive:
            styles = ET.fromstring(archive.read('xl/styles.xml'))
            fills = styles.findall('s:dxfs/s:dxf/s:fill/s:patternFill', ns)
            self.assertEqual([fill.find('s:fgColor', ns).get('rgb') for fill in fills],
                             ['FFFFFFFF', 'FFDCE6F1', 'FFFFFFFF', 'FFEDF3F8'])
            self.assertTrue(all(fill.find('s:bgColor', ns).get('rgb') ==
                                fill.find('s:fgColor', ns).get('rgb') for fill in fills))
        self.assertEqual(self.snapshot(), before)

    def test_filename_opens_database_without_injected_connection(self):
        self.assertEqual(server.supplier_registry_excel_filename({'framework_id': ['f1']}),
                         '03110000-1_pqm-suppliers.xlsx')

    def test_multi_framework_endpoint_keeps_generic_filename(self):
        status, headers, raw = self.get('registry_only=1&framework_id=f1&framework_id=f2')
        self.assertEqual(status, 200)
        self.assertEqual(headers['Content-Disposition'], 'attachment; filename="pqm-suppliers.xlsx"')
        self.assertEqual(load_workbook(io.BytesIO(raw)).active.max_row, 3)

    def test_repeated_export_is_read_only(self):
        before = self.snapshot()
        for _ in range(2):
            self.assertEqual(self.get('registry_only=1&framework_id=f1')[0], 200)
        self.assertEqual(self.snapshot(), before)


if __name__ == '__main__':
    unittest.main()
