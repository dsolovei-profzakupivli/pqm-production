"""Presentation regressions; all inputs are synthetic, with no runtime access."""
import io
import unittest
from datetime import datetime
from openpyxl import load_workbook
import supplier_xlsx
import test_supplier_filter_excel as fixtures


class XlsxFormattingTests(unittest.TestCase):
    def rows(self):
        return [{'values':[code,name,stamp,'Активний',email,phone]} for code,name,stamp,email,phone in [
            ('0001','=literal\r\nУкраїна','2026-10-14T23:30:00Z','Case@EXAMPLE.invalid','+380001; 001'),
            ('0002','Ранній','2026-10-15T02:30:00.004491+03:00','',''),
            ('0003','Пізніший','2026-10-15T02:30:00.004492+03:00','',''),
            ('0004','Немає дати','','',''),('0005','Некоректна','invalid','','')]]

    def sheet(self,rows=None):
        return load_workbook(io.BytesIO(supplier_xlsx.workbook(self.rows() if rows is None else rows))).active

    def test_exact_order_then_invalid_dates_last_without_input_mutation(self):
        rows=self.rows();original=repr(rows);ws=self.sheet(rows)
        self.assertEqual([ws.cell(i,1).value for i in range(2,7)],['0003','0002','0001','0004','0005'])
        self.assertEqual(repr(rows),original)
        self.assertIsNone(ws['C5'].value);self.assertIsNone(ws['C6'].value)

    def test_typed_kyiv_calendar_date_and_format(self):
        ws=self.sheet()
        for i in range(2,5):
            c=ws.cell(i,3)
            self.assertEqual(c.value,datetime(2026,10,15));self.assertEqual(c.data_type,'d')
            self.assertEqual(c.number_format,'dd.mm.yyyy')

    def test_dst_date_boundaries(self):
        rows=[{'values':['01','Весна','2026-03-28T22:30:00Z','','','']},
              {'values':['02','Осінь','2026-10-24T21:30:00Z','','','']}]
        ws=self.sheet(rows)
        self.assertEqual(ws['C2'].value,datetime(2026,10,25))
        self.assertEqual(ws['C3'].value,datetime(2026,3,29))

    def test_literal_values_formula_guard_crlf_and_contacts(self):
        ws=self.sheet();self.assertEqual(ws['A4'].value,'0001');self.assertEqual(ws['A4'].data_type,'s')
        self.assertEqual(ws['B4'].value,'=literal\r\nУкраїна');self.assertEqual(ws['B4'].data_type,'s')
        self.assertEqual(ws['E4'].value,'Case@EXAMPLE.invalid');self.assertEqual(ws['F4'].value,'+380001; 001')
        self.assertTrue(all(c.data_type!='f' for row in ws for c in row))

    def test_table_filter_freeze_borders_and_no_protection(self):
        ws=self.sheet();self.assertEqual(ws.max_column,6)
        self.assertEqual([c.value for c in ws[1]],list(supplier_xlsx.HEADERS))
        self.assertEqual(ws.tables['PqmSuppliers'].ref,'A1:F6')
        self.assertEqual(ws.auto_filter.ref,'A1:F6');self.assertEqual(ws.freeze_panes,'A2')
        self.assertFalse(ws.protection.sheet);self.assertEqual(ws['A2'].border.bottom.style,'thin')

    def test_top_alignment_wrapping_and_height(self):
        ws=self.sheet()
        for row in ws:
            for c in row:self.assertEqual(c.alignment.vertical,'top')
        for column in 'BEF':self.assertTrue(ws[column+'2'].alignment.wrap_text)
        rows=[{'values':['01','Довга назва '*30,'2026-10-14T14:00:00+03:00','','Email'*50,'+380 '*30]}]
        self.assertGreater(self.sheet(rows).row_dimensions[2].height,30)

    def test_empty_table_and_missing_fields(self):
        ws=self.sheet([]);self.assertEqual(ws.max_row,1);self.assertEqual(ws.tables['PqmSuppliers'].ref,'A1:F1')

    def test_no_silent_row_height_clipping(self):
        with self.assertRaisesRegex(ValueError,'висоту'):
            self.sheet([{'values':['1','X'*32767,'','','','']}])


class FilenameTests(unittest.TestCase):
    def test_framework_search_name_category_and_intersection(self):
        fixture=fixtures.SupplierFilterExcelTests();fixture.setUp()
        try:
            import server
            fixture.con.execute("INSERT INTO framework_service_directory(pretty_id,framework_id,category,synced_at) VALUES('UA-F-f1','f1','Металеві вироби','synthetic')")
            for search in ('ВІДБІР F1','МЕТАЛЕВІ ВИРОБИ','UA-F-f1','03110000-1'):
                result=server.framework_registry({'search':[search],'dk_code':['03110000-1']},fixture.con)
                self.assertIn('f1',[item['framework_id'] for item in result['items']])
            self.assertEqual(server.framework_registry({'search':['Металеві'],'dk_code':['03220000-2']},fixture.con)['total'],0)
        finally:fixture.doCleanups()

    def test_filename_uses_cached_framework_not_search_input(self):
        fixture=fixtures.SupplierFilterExcelTests();fixture.setUp()
        try:
            import server
            params={'registry_only':['1'],'framework_id':['f1'],'dk_code':['FAKE']}
            self.assertEqual(server.supplier_registry_excel_filename(params,connection=fixture.con),'03110000-1_pqm-suppliers.xlsx')
            params['framework_id']=['f1','f2']
            self.assertEqual(server.supplier_registry_excel_filename(params,connection=fixture.con),'pqm-suppliers.xlsx')
        finally:fixture.doCleanups()


if __name__=='__main__':unittest.main()
