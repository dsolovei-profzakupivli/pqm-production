"""Synthetic in-memory DB and in-memory XLSX. No sync, network or runtime writes."""
import io
import json
import sqlite3
import unittest
import warnings
import re
import subprocess
from html.parser import HTMLParser
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import server
import reference_directories
import supplier_selection
import supplier_xlsx
import auth_access
from datetime import datetime


class SupplierFilterExcelTests(unittest.TestCase):
    def test_framework_search_discards_stale_success_and_error_without_sync(self):
        result=subprocess.run(['node','validation/framework_registry_race_ui.cjs'],
          cwd=Path(__file__).parent,capture_output=True,text=True,timeout=20)
        self.assertEqual(result.returncode,0,result.stderr)

    def setUp(self):
        self.con=sqlite3.connect(':memory:');self.con.row_factory=sqlite3.Row
        self.con.create_function('DIGITS',1,lambda v:''.join(c for c in str(v or '') if c.isdigit()))
        self.con.create_function('CASEFOLD',1,lambda v:str(v or '').casefold())
        server.edr_sync_v2.register_verification_sql_functions(self.con)
        with patch('reference_directories.sqlite3.connect',return_value=self.con):
            reference_directories.init_reference_tables(':memory:')
        with patch.object(server,'db',return_value=self.con),patch.object(server,'DB_PATH',Path('memory-only-fixture.sqlite3')):
            server.init_db()
        self.addCleanup(self.con.close)
        for fid,dk in [('f1','03110000-1'),('f2','03220000-2'),('f3','03110000-1')]:
            self.con.execute('INSERT INTO frameworks(id,pretty_id,title,dk_code,status,raw_json,synced_at) VALUES(?,?,?,?,?,?,?)',
              (fid,'UA-F-'+fid,'Відбір '+fid,dk,'active',json.dumps({'qualificationPeriod':{'endDate':'2099-01-01T00:00:00+02:00'}}),'fixture'))

    def add(self,sid,code='0000000123',fid='f1',q='active',contract='active',day=1,email=None,phone='+380001234567',name='Постачальник «Тест»'):
        qid='q'+sid;date=f'2026-10-{day:02d}T12:00:00+03:00'
        raw=json.dumps({'tenderers':[{'contactPoint':{'email':email if email is not None else sid+'@example.invalid','telephone':phone}}]})
        self.con.execute('INSERT INTO submissions(id,framework_id,supplier_code,supplier_name,date_published,status,qualification_id,raw_json,synced_at) VALUES(?,?,?,?,?,?,?,?,?)',
          (sid,fid,code,name,date,'complete',qid,raw,'2099-01-01'))
        self.con.execute('INSERT INTO qualifications(id,framework_id,submission_id,status,decision_date,raw_json,synced_at) VALUES(?,?,?,?,?,?,?)',
          (qid,fid,sid,q,'2099-01-01',json.dumps({'id':qid,'status':q,'date':date,'submissionID':sid}),'2099-01-01'))
        if contract:
            self.con.execute('INSERT INTO registry_contracts(id,framework_id,qualification_id,supplier_code,status,raw_json,synced_at) VALUES(?,?,?,?,?,?,?)',
              ('r'+sid,fid,qid,code,contract,json.dumps({'date':date}),'fixture'))
        self.con.execute('INSERT OR IGNORE INTO supplier_registry_summary(supplier_code,supplier_name,refreshed_at) VALUES(?,?,?)',(code,name,'fixture'))

    def listing(self,**params):
        return server.list_qualified_suppliers({k:[str(v)] for k,v in params.items()},connection=self.con)

    def export(self,**params):
        return server.supplier_registry_excel({k:[str(v)] for k,v in params.items()},connection=self.con)

    def values(self,raw):
        with ZipFile(io.BytesIO(raw)) as z:
            root=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        ns={'s':supplier_xlsx.NS}
        self.assertEqual(root.findall('.//s:f',ns),[])
        from openpyxl import load_workbook
        sheet=load_workbook(io.BytesIO(raw),read_only=True).active
        return [[cell.value if cell.value is not None else '' for cell in row] for row in sheet]

    def test_intersection_filters_before_aggregation(self):
        self.add('a',fid='f1');self.add('b',fid='f2',contract='terminated')
        result=self.listing(framework_id='f2',status='terminated')
        self.assertEqual(result['total'],1)
        self.assertEqual(result['items'][0]['active_count'],0)
        self.assertEqual(result['items'][0]['prozorro_status'],'Неактивний')
        self.assertEqual(self.listing(dk_code='03110000-1',framework_id='f2')['total'],0)
        self.assertEqual(self.listing(framework_id='f1',status='terminated')['total'],0)
        self.assertEqual(self.listing(dk_code='03110000-1',status='active')['total'],1)

    def test_registry_filter_does_not_require_application_backlink(self):
        self.add('a')
        self.con.execute("UPDATE submissions SET qualification_id=NULL WHERE id='a'")
        row=self.listing(framework_id='f1',status='active')['items'][0]
        self.assertEqual(row['code'],'0000000123')
        self.assertEqual(row['active_count'],1)
        self.assertEqual(self.listing(framework_id='f2',status='active')['total'],0)

    def test_registry_only_membership_keeps_dk_framework_status_intersection(self):
        self.con.execute("INSERT INTO registry_contracts(id,framework_id,supplier_code,status,raw_json,synced_at) VALUES('orphan','f1','0000000456','active','{}','fixture')")
        # The registry endpoint remains fail-closed until its summary is built.
        self.con.execute("INSERT INTO supplier_registry_summary(supplier_code,supplier_name,refreshed_at) VALUES('0000000456','Registry-only supplier','fixture')")
        self.assertEqual(self.listing(framework_id='f1',dk_code='03110000-1',status='active')['total'],1)
        self.assertEqual(self.listing(framework_id='f1',dk_code='03220000-2',status='active')['total'],0)
        self.assertEqual(self.listing(framework_id='f2',status='active')['total'],0)
        self.assertEqual(self.listing(framework_id='f1',status='terminated')['total'],0)
        self.assertEqual(self.listing(framework_id='f1',status='not_registered')['total'],0)

    def test_current_admission_wins_newer_pending_and_contacts_stay_linked(self):
        self.add('admit',day=1);self.add('pending',day=2,q='pending',contract=None)
        result=self.values(self.export(framework_id='f1'))
        self.assertEqual(len(result),2)
        self.assertEqual(result[1][2],datetime(2026,10,1))
        self.assertEqual(result[1][3],'Активний')
        self.assertEqual(result[1][4],'admit@example.invalid')

    def test_terminated_admission_not_replaced_by_pending(self):
        self.add('old',contract='terminated');self.add('new',day=2,q='pending',contract=None,email='',phone='')
        result=self.values(self.export(framework_id='f1'))[1]
        self.assertEqual(result[2],datetime(2026,10,1));self.assertEqual(result[3],'Неактивний')
        self.assertEqual(result[4],'old@example.invalid')

    def test_admission_not_replaced_by_rejected(self):
        self.add('old',contract='terminated');self.add('rejected',day=3,q='unsuccessful',contract=None)
        result=self.values(self.export(framework_id='f1'))[1]
        self.assertEqual(result[2],datetime(2026,10,1))
        self.assertEqual(result[3],'Неактивний');self.assertEqual(result[4],'old@example.invalid')

    def test_one_row_per_supplier_latest_admission_across_frameworks(self):
        self.add('one',fid='f1');self.add('two',fid='f2',day=2)
        self.assertEqual(len(self.values(self.export())),2)
        self.assertEqual(self.values(self.export())[1][4],'two@example.invalid')
        self.assertEqual(self.values(self.export(framework_id='f1'))[1][4],'one@example.invalid')

    def test_all_pages_exported_without_duplicates(self):
        for i in range(23):self.add(str(i),code=f'{i:010d}')
        self.assertEqual(len(self.listing(framework_id='f1',size=10)['items']),10)
        rows=self.values(self.export(framework_id='f1',page=3,size=10))
        self.assertEqual(len(rows),24);self.assertEqual(len({r[0] for r in rows[1:]}),23)

    def test_unicode_leading_zeros_phone_and_formula_injection_are_text(self):
        self.add('a',name='=HYPERLINK("bad")',email='@unsafe',phone='+380001234567')
        raw=self.export(framework_id='f1');rows=self.values(raw)
        self.assertEqual(rows[1][0],'0000000123');self.assertEqual(rows[1][1],'=HYPERLINK("bad")')
        self.assertEqual(rows[1][4],'@unsafe');self.assertEqual(rows[1][5],'+380001234567')
        from openpyxl import load_workbook
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            sheet=load_workbook(io.BytesIO(raw),read_only=True).active
            self.assertEqual(sheet['B2'].data_type,'s');self.assertEqual(sheet['F2'].value,'+380001234567')

    def test_readonly_no_mutations_and_prod_sandbox_parity(self):
        self.add('a');self.con.commit();self.con.execute('PRAGMA query_only=ON')
        before=self.con.total_changes
        with patch.object(server,'SANDBOX_MODE',True):one=self.values(self.export(framework_id='f1'))
        with patch.object(server,'SANDBOX_MODE',False):two=self.values(self.export(framework_id='f1'))
        self.assertEqual(one,two);self.assertEqual(self.con.total_changes,before)

    def test_unknown_framework_empty_export_and_parameter_injection(self):
        self.add('a')
        self.assertEqual(self.listing(framework_id="f1' OR 1=1 --")['total'],0)
        self.assertEqual(len(self.values(self.export(framework_id='missing'))),1)

    def test_suspended_and_expired_admission(self):
        self.add('a',contract='suspended')
        self.assertEqual(self.listing(status='suspended')['total'],1)
        self.assertEqual(self.values(self.export(status='suspended'))[1][3],'Призупинений')
        self.con.execute("UPDATE frameworks SET status='complete' WHERE id='f1'")
        self.add('pending',q='pending',contract=None,day=2)
        self.assertEqual(self.values(self.export(framework_id='f1'))[1][3],'Призупинений')

    def test_pending_rejected_only_excluded_and_empty_contacts_preserved(self):
        self.add('p',code='1',q='pending',contract=None)
        self.add('r',code='2',q='unsuccessful',contract=None)
        self.add('a',code='3',email='',phone='')
        rows=self.values(self.export())
        self.assertEqual(len(rows),2);self.assertEqual(rows[1][0],'3')
        self.assertEqual(rows[1][4:],['','']);self.assertEqual(len(rows[0]),6)

    def test_multiple_selected_frameworks_latest_not_active_priority(self):
        self.add('a',fid='f1',day=1)
        self.add('b',fid='f2',day=2,contract='terminated')
        self.add('outside',fid='f3',day=3)
        raw=server.supplier_registry_excel({'framework_id':['f1','f2']},connection=self.con)
        rows=self.values(raw)
        self.assertEqual(len(rows),2);self.assertEqual(rows[1][4],'b@example.invalid')
        self.assertEqual(rows[1][3],'Неактивний')
        raw=server.supplier_registry_excel({'framework_id':['f1','f2'],'status':['active']},connection=self.con)
        self.assertEqual(len(self.values(raw)),1)

    def test_4453_synthetic_active_inactive_coverage(self):
        self.con.execute("UPDATE frameworks SET pretty_id='4453' WHERE id='f1'")
        self.add('active',code='1');self.add('inactive',code='2',contract='terminated')
        self.add('pending',code='3',q='pending',contract=None)
        rows=self.values(self.export(framework_id='f1'))
        self.assertEqual({r[0] for r in rows[1:]},{'1','2'})
        self.assertEqual({r[3] for r in rows[1:]},{'Активний','Неактивний'})

    def test_admission_order_uses_exact_instants_not_submission_or_sync(self):
        self.add('a',fid='f1',day=20);self.add('b',fid='f2',day=1,contract='terminated')
        for sid,stamp in [('a','2026-10-01T12:00:00.004491+03:00'),('b','2026-10-01T09:00:00.004492Z')]:
            raw=json.loads(self.con.execute('SELECT raw_json FROM qualifications WHERE id=?',('q'+sid,)).fetchone()[0])
            raw['date']=stamp
            self.con.execute('UPDATE qualifications SET raw_json=? WHERE id=?',(json.dumps(raw),'q'+sid))
        row=self.values(self.export())[1]
        self.assertEqual(row[2],datetime(2026,10,1))
        selected=supplier_selection.export_rows(self.con,{}, {'0000000123'})
        self.assertEqual(selected[0]['values'][2],'2026-10-01T09:00:00.004492Z')
        self.assertEqual(row[4],'b@example.invalid')

    def test_invalid_admission_timestamp_cannot_fall_back_to_submission(self):
        self.add('a')
        self.con.execute("UPDATE qualifications SET raw_json=json_set(raw_json,'$.date','invalid')")
        self.assertEqual(len(self.values(self.export())),1)

    def registry_export(self,ids):
        return self.values(server.supplier_registry_excel({'framework_id':ids,'registry_only':['1']},connection=self.con))

    def test_framework_table_unique_counts_and_no_bids_dependency(self):
        self.add('a',code='1');self.add('b',code='2',contract='suspended');self.add('c',code='3',contract='terminated')
        self.add('pending',code='4',q='pending',contract=None)
        self.add('again',code='1',day=2)
        self.con.commit();self.con.execute('PRAGMA query_only=ON');before=self.con.total_changes
        queries=[];self.con.set_trace_callback(queries.append)
        data=server.framework_registry({'search':['f1']},self.con)
        c=data['items'][0]['counts']
        self.assertEqual((c['total'],c['active'],c['suspended'],c['terminated'],c['other']),(3,1,1,1,0))
        self.assertEqual(c['total'],c['active']+c['suspended']+c['terminated'])
        self.assertFalse(any(re.search(r'\b(bids|tenders|agreements)\b',q,re.I) for q in queries))
        self.assertEqual(self.con.total_changes,before)

    def test_block90_and_completed_ban_follow_current_contract_status(self):
        self.add('blocked',code='44368854',contract='suspended')
        self.add('resumed',code='1',contract='active')
        for sid,state in [('blocked','scheduled'),('resumed','met')]:
            milestone=[{'type':'ban','status':state,'dateModified':'2026-09-03T12:00:00+03:00','dueDate':'2026-12-03T12:00:00+02:00'}]
            self.con.execute('UPDATE registry_contracts SET milestones_json=? WHERE id=?',(json.dumps(milestone),'r'+sid))
        c=supplier_selection.registry_counts(self.con)['f1']
        self.assertEqual((c['total'],c['active'],c['suspended'],c['terminated']),(2,1,1,0))
        rows=self.registry_export(['f1'])
        self.assertEqual({r[0]:r[3] for r in rows[1:]},{'1':'Активний','44368854':'Призупинений'})
        self.assertEqual(rows[2][2],datetime(2026,10,1))

    def test_registry_single_multiple_and_pagination_complete(self):
        self.add('old',fid='f1',day=1);self.add('new',fid='f2',day=2,contract='terminated')
        self.add('pending',fid='f2',day=3,q='pending',contract=None)
        self.assertEqual(self.registry_export(['f1'])[1][4],'old@example.invalid')
        rows=self.registry_export(['f1','f2'])
        self.assertEqual(len(rows),2);self.assertEqual(rows[1][4],'new@example.invalid')
        self.assertEqual(rows[1][3],'Неактивний');self.assertEqual(len(rows[1]),6)
        for i in range(23):self.add(str(i),code=f'{i:010d}')
        rows=self.values(server.supplier_registry_excel({'registry_only':['1'],'framework_id':['f1'],'page':['99'],'size':['1']},connection=self.con))
        self.assertEqual(len(rows),25)

    def test_missing_admission_date_not_replaced_by_contract_date(self):
        self.add('old',contract='terminated')
        self.con.execute("UPDATE qualifications SET status='cancelled',raw_json=json_set(raw_json,'$.status','cancelled','$.date','2026-10-20T00:00:00Z')")
        with self.assertRaisesRegex(ValueError,'Не підтверджено дату допуску'):
            self.registry_export(['f1'])

    def test_registry_formula_injection_and_missing_contacts(self):
        self.add('a',name='=HYPERLINK("bad")',email='',phone='')
        row=self.registry_export(['f1'])[1]
        self.assertEqual(row[0],'0000000123');self.assertEqual(row[1],'=HYPERLINK("bad")')
        self.assertEqual(row[4:],['',''])

    def test_equal_admission_conflicting_states_fail_closed(self):
        self.add('a');self.add('b',contract='terminated')
        self.assertEqual(supplier_selection.registry_counts(self.con)['f1']['other'],1)
        with self.assertRaisesRegex(ValueError,'Неоднозначний поточний статус'):
            self.registry_export(['f1'])

    def test_registry_unknown_framework_cannot_silently_export_empty(self):
        with self.assertRaisesRegex(ValueError,'відсутній у локальному кеші'):
            self.registry_export(['missing'])

    def test_registry_export_has_no_unfinished_edr_dependency(self):
        self.add('a')
        import sys
        with patch.dict(sys.modules,{'edr_admission_read_model':None}):
            self.assertEqual(len(self.registry_export(['f1'])),2)

    def test_supplier_2746305948_original_crlf_roundtrip(self):
        from openpyxl import load_workbook
        original='ФОП "ПРОДАН НАТАЛІЯ ВІКТОРІВНА\r\n"'
        self.add('crlf',code='2746305948',name=original)
        payload=server.supplier_registry_excel({'framework_id':['f1'],'registry_only':['1']},connection=self.con)
        with warnings.catch_warnings():
            warnings.simplefilter('ignore');ws=load_workbook(io.BytesIO(payload),read_only=True).active
        self.assertEqual(ws['B2'].value,original);self.assertEqual(ws['B2'].data_type,'s')

    def test_unknown_status_explicit_report_and_export_fail_closed(self):
        self.add('a',contract='unexpected')
        c=supplier_selection.registry_counts(self.con)['f1']
        self.assertEqual((c['total'],c['other']),(1,1));self.assertEqual(c['issues'][0]['statuses'],['unexpected'])
        with self.assertRaisesRegex(ValueError,'Невідомий статус'):self.registry_export(['f1'])

    def test_registry_sort_filter_and_pagination(self):
        self.add('a');self.add('b',fid='f2',code='2');self.add('c',fid='f2',code='3')
        data=server.framework_registry({'sort':['supplier_count'],'direction':['desc'],'size':['1']},self.con)
        self.assertEqual(data['items'][0]['framework_id'],'f2');self.assertEqual(data['pages'],3)
        data=server.framework_registry({'dk_code':['03110000-1'],'search':['f1']},self.con)
        self.assertEqual([i['framework_id'] for i in data['items']],['f1'])
        self.assertEqual(auth_access.permission_key('GET','/api/framework-registry'),'frameworks.read')

    def test_framework_ui_separates_bids_and_scopes_downloads(self):
        root=Path(__file__).parent;js=(root/'app.js').read_text(encoding='utf-8')
        main=js.split('async function loadFrameworkAnalytics(){',1)[1].split('async function loadFrameworkBidsSection()',1)[0]
        self.assertIn('/framework-registry?',main);self.assertNotIn('/framework-analytics?',main)
        self.assertIn('data-registry-framework',main);self.assertIn('registryExportUrl([f.id])',main)
        script=js.split('function registryExportUrl(ids){',1)[1].split('\n',1)[0]
        result=subprocess.run(['node','-e',"const API='/api';function registryExportUrl(ids){"+script+";const u=registryExportUrl(['f1','f2']);const q=new URL(u,'https://local.test').searchParams;if(q.getAll('framework_id').join(',')!=='f1,f2'||q.get('registry_only')!=='1')process.exit(1)"],capture_output=True,text=True,timeout=10)
        self.assertEqual(result.returncode,0,result.stderr)

    def test_export_requires_supplier_read_permission(self):
        self.assertEqual(auth_access.permission_key('GET','/api/suppliers-registry/export.xlsx'),'suppliers.read')

    def test_ui_visibility_and_shared_filter_contract(self):
        class Visibility(HTMLParser):
            def __init__(self):super().__init__();self.stack=[];self.found={}
            def handle_starttag(self,tag,attrs):
                attrs=dict(attrs)
                if attrs.get('id')=='suppliersView':attrs.pop('hidden',None)
                hidden=('hidden' in attrs or attrs.get('aria-hidden')=='true' or any(self.stack))
                if attrs.get('id'):self.found.setdefault(attrs['id'],[]).append(hidden)
                if tag not in {'input','br','hr','img','meta','link','source','wbr'}:self.stack.append(hidden)
            def handle_endtag(self,tag):
                if self.stack:self.stack.pop()
        root=Path(__file__).parent
        html=(root/'index.html').read_text(encoding='utf-8')
        supplier_view=re.search(r'<main id="suppliersView".*?</main>',html,re.S).group()
        parser=Visibility();parser.feed(supplier_view)
        for key in ['supplierRegistryDkFilter','supplierRegistryFrameworkFilter']:
            self.assertEqual(parser.found[key],[False],key)
        self.assertEqual(parser.found['supplierRegistryFramework'],[True])
        self.assertNotIn('supplierRegistryExcel',parser.found)
        self.assertEqual(parser.found['supplierRegistryEdrStatus'],[True])
        js=(root/'app.js').read_text(encoding='utf-8')
        self.assertIn('/suppliers-registry/export.xlsx?',js)
        function=re.search(r'function supplierRegistryFilterParams\(\)\{.*?\n\}',js,re.S).group()
        script="const $=id=>({value:id.includes('Dk')?'03110000-1':' Проба ',selectedOptions:[{value:'f1'},{value:'f2'}]});const supplierRegistryRisk='';"+function+";const q=supplierRegistryFilterParams();if(q.getAll('framework_id').join(',')!=='f1,f2'||q.get('dk_code')!=='03110000-1'||q.has('page')||q.has('size'))process.exit(1);"
        result=subprocess.run(['node','-e',script],capture_output=True,text=True,timeout=10)
        self.assertEqual(result.returncode,0,result.stderr)

    def test_excel_rejects_lossy_or_invalid_cells(self):
        for value in ['bad\x00value','x'*32768]:
            with self.assertRaises(ValueError):supplier_xlsx.workbook([{'values':[value,'','','','','']}])


if __name__=='__main__':unittest.main()
