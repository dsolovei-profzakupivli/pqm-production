"""Isolated production-baseline schema; synthetic rows only, no runtime DB."""
import io,json,sqlite3,unittest,subprocess
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
import server,reference_directories,auth_access


class SupplierFilterExcelTests(unittest.TestCase):
    """Shared fixture (no test methods): schema comes from exact PROD source."""
    def setUp(self):
        self.con=sqlite3.connect(':memory:');self.con.row_factory=sqlite3.Row
        self.con.create_function('DIGITS',1,lambda v:''.join(c for c in str(v or '') if c.isdigit()))
        self.con.create_function('CASEFOLD',1,lambda v:str(v or '').casefold())
        server.edr_sync_v2.register_verification_sql_functions(self.con)
        with patch('reference_directories.sqlite3.connect',return_value=self.con):reference_directories.init_reference_tables(':memory:')
        with patch.object(server,'db',return_value=self.con),patch.object(server,'DB_PATH',Path('memory-only-fixture.sqlite3')):server.init_db()
        self.addCleanup(self.con.close)
        for fid,dk in [('f1','03110000-1'),('f2','03220000-2'),('f3','03110000-1')]:
            self.con.execute('INSERT INTO frameworks(id,pretty_id,title,dk_code,status,raw_json,synced_at) VALUES(?,?,?,?,?,?,?)',
              (fid,'UA-F-'+fid,'Відбір '+fid,dk,'active',json.dumps({'qualificationPeriod':{'endDate':'2099-01-01T00:00:00+02:00'}}),'fixture'))

    def add(self,sid,code='0000000123',fid='f1',q='active',contract='active',day=1,email=None,phone='+380001234567',name='Постачальник «Тест»'):
        qid='q'+sid;date=f'2026-10-{day:02d}T12:00:00+03:00'
        raw=json.dumps({'tenderers':[{'contactPoint':{'email':email if email is not None else sid+'@example.invalid','telephone':phone}}]})
        self.con.execute('INSERT INTO submissions(id,framework_id,supplier_code,supplier_name,date_published,status,qualification_id,raw_json,synced_at) VALUES(?,?,?,?,?,?,?,?,?)',
          (sid,fid,code,name,date,'complete',qid,raw,'2099-01-01'))
        self.con.execute('INSERT INTO qualifications(id,framework_id,submission_id,status,decision_date,raw_json,synced_at) VALUES(?,?,?,?,?,?,?)',
          (qid,fid,sid,q,'2099-01-01',json.dumps({'id':qid,'status':q,'date':date,'submissionID':sid}),'2099-01-01'))
        if contract:
            self.con.execute('INSERT INTO registry_contracts(id,framework_id,qualification_id,supplier_code,status,raw_json,synced_at) VALUES(?,?,?,?,?,?,?)',
              ('r'+sid,fid,qid,code,contract,json.dumps({'date':date}),'fixture'))
        self.con.execute('INSERT OR IGNORE INTO supplier_registry_summary(supplier_code,supplier_name,refreshed_at) VALUES(?,?,?)',(code,name,'fixture'))

    def export(self,ids):
        return load_workbook(io.BytesIO(server.supplier_registry_excel({'registry_only':['1'],'framework_id':ids},connection=self.con))).active


class ProductionFrameworkTests(SupplierFilterExcelTests):
    def test_late_search_response_and_error_cannot_replace_new_results(self):
        result=subprocess.run(['node','validation/framework_registry_race_ui.cjs'],
          cwd=Path(__file__).parent,capture_output=True,text=True,timeout=20)
        self.assertEqual(result.returncode,0,result.stderr)

    def test_counts_statuses_and_complete_export(self):
        self.add('active',code='1');self.add('suspended',code='2',contract='suspended');self.add('terminated',code='3',contract='terminated')
        data=server.framework_registry({'search':['f1']},self.con);c=data['items'][0]['counts']
        self.assertEqual([c[k] for k in ('total','active','suspended','terminated')],[3,1,1,1])
        ws=self.export(['f1']);self.assertEqual(ws.max_row,4)
        self.assertEqual({ws.cell(i,4).value for i in range(2,5)},{'Активний','Призупинений','Неактивний'})

    def test_multiple_framework_dedup_and_contacts_from_chosen_admission(self):
        self.add('old',day=1);self.add('new',fid='f2',day=2,contract='terminated')
        self.add('pending',fid='f2',q='pending',contract=None,day=3)
        ws=self.export(['f1','f2']);self.assertEqual(ws.max_row,2)
        self.assertEqual(ws['E2'].value,'new@example.invalid');self.assertEqual(ws['C2'].value.day,2)

    def test_pending_and_rejected_without_admission_not_included(self):
        self.add('pending',q='pending',contract=None);self.add('rejected',code='2',q='unsuccessful',contract=None)
        self.assertEqual(self.export(['f1']).max_row,1)

    def test_query_does_not_mutate_schema_settings_or_data(self):
        self.add('one');self.con.commit();self.con.execute('PRAGMA query_only=ON')
        schema=list(self.con.execute('SELECT sql FROM sqlite_master ORDER BY name'));changes=self.con.total_changes
        server.framework_registry({},self.con);self.export(['f1'])
        self.assertEqual(list(self.con.execute('SELECT sql FROM sqlite_master ORDER BY name')),schema)
        self.assertEqual(self.con.total_changes,changes)

    def test_unknown_scope_or_status_fail_closed(self):
        with self.assertRaises(ValueError):self.export(['missing'])
        self.add('one',contract='unexpected')
        with self.assertRaisesRegex(ValueError,'статус'):self.export(['f1'])

    def test_no_silent_admission_timestamp_substitution(self):
        self.add('one');self.con.execute("UPDATE qualifications SET raw_json=json_set(raw_json,'$.date','invalid')")
        with self.assertRaises(ValueError):self.export(['f1'])

    def test_sort_search_and_pagination(self):
        self.add('one');self.add('two',fid='f2',code='2');self.add('three',fid='f2',code='3')
        data=server.framework_registry({'sort':['supplier_count'],'direction':['desc'],'size':['1']},self.con)
        self.assertEqual(data['items'][0]['framework_id'],'f2');self.assertEqual(data['pages'],3)
        self.assertEqual(server.framework_registry({'search':['ВІДБІР F1'],'dk_code':['03220000-2']},self.con)['total'],0)

    def test_existing_rbac_boundaries(self):
        self.assertEqual(auth_access.permission_key('GET','/api/framework-registry'),'frameworks.read')
        self.assertEqual(auth_access.permission_key('GET','/api/suppliers-registry/export.xlsx'),'suppliers.read')
