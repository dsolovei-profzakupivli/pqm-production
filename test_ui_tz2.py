"""TZ2: disposable SQLite/DOCX fixtures only, never runtime data."""
from contextlib import closing
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch
from datetime import datetime
import json,sqlite3,tempfile,threading,unittest,subprocess
from zipfile import ZipFile
import audit_reader
import violation_protocol_summary as protocols
import server
import test_ui_improvements as counts_fixture
import test_audit_reader as audit_fixture

class AuditTz2Tests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.path=Path(self.tmp.name)/'audit.sqlite3'
        with closing(sqlite3.connect(self.path)) as con:
            con.execute('CREATE TABLE audit_log(id INTEGER PRIMARY KEY,submission_id TEXT,changed_at TEXT,changed_by TEXT,field_name TEXT,old_value TEXT,new_value TEXT)')
            con.executemany('INSERT INTO audit_log VALUES(?,?,?,?,?,?,?)',[(i,'s'+str(i),'2026-10-17T06:00:00Z','УО','remarks','абґ '*300,'Не виконано '+str(i)) for i in range(1200)])
            con.commit()
    def test_exact_count_unicode_empty_and_out_of_range(self):
        for query,total in [({'search':['НЕ ВИКОНАНО']},1200),({'search':['відсутній']},0),({'search':['s1199']},1)]:
            data=audit_reader.read(self.path,query);self.assertEqual(data['total'],total)
        data=audit_reader.read(self.path,{'search':['НЕ ВИКОНАНО'],'page':['999']})
        self.assertEqual(data['total'],1200);self.assertEqual(data['items'],[])
    def test_deadline_retained_and_controlled(self):
        ticks=[0]
        def clock():ticks[0]+=1;return ticks[0]*2
        with patch.object(audit_reader.time,'monotonic',clock):
            with self.assertRaises(audit_reader.AuditResourceError) as error:audit_reader.read(self.path,{'search':['відсутній']})
        self.assertEqual(error.exception.status,503)
    def test_unicode_match_is_evaluated_once_per_record_for_count_and_page(self):
        real=sqlite3.connect;calls=[0]
        class Counting(sqlite3.Connection):
            def create_function(self,name,arity,fn,**kw):
                def wrapped(*args):calls[0]+=1;return fn(*args)
                return super().create_function(name,arity,wrapped,**kw)
        with patch.object(audit_reader.sqlite3,'connect',lambda *a,**kw:real(*a,**kw,factory=Counting)):
            result=audit_reader.read(self.path,{'search':['не виконано'],'page':['2']})
        self.assertEqual(result['total'],1200);self.assertEqual(len(result['items']),100);self.assertEqual(calls[0],1200)
    def test_cancellation_is_scoped_and_stale_request_cannot_restart(self):
        client='test-'+Path(self.tmp.name).name
        event=threading.Event();release=threading.Event();real=sqlite3.connect
        class Slow(sqlite3.Connection):
            def create_function(self,name,arity,fn,**kw):
                def wrapped(*args):
                    if args[0]=='hold':event.set();release.wait(3)
                    return fn(*args)
                return super().create_function(name,arity,wrapped,**kw)
        def connect(*a,**kw):return real(*a,**kw,factory=Slow)
        params={'search':['hold'],'client_id':[client],'request_id':['1']}
        with patch.object(audit_reader.sqlite3,'connect',connect),ThreadPoolExecutor(2) as pool:
            first=pool.submit(audit_reader.read,self.path,params,actor='alice')
            self.assertTrue(event.wait(2))
            # Bob is not blocked by Alice's active search or cancellation.
            other=audit_reader.read(self.path,{'client_id':[client],'request_id':['1']},actor='bob')
            self.assertEqual(other['total'],1200)
            self.assertTrue(audit_reader.read(self.path,{'client_id':[client],'request_id':['2'],'cancel':['1']},actor='alice')['cancelled'])
            release.set()
            with self.assertRaises(audit_reader.AuditSuperseded):first.result()
        with self.assertRaises(audit_reader.AuditSuperseded):audit_reader.read(self.path,params,actor='alice')
        self.assertEqual(audit_reader.read(self.path,{'client_id':[client],'request_id':['3']},actor='alice')['total'],1200)
    def test_long_cyrillic_data_concurrent_readers_and_no_schema_writes(self):
        with closing(sqlite3.connect(self.path)) as con:before=list(con.iterdump())
        with ThreadPoolExecutor(4) as pool:
            results=list(pool.map(lambda n:audit_reader.read(self.path,{'search':['не виконано']},actor=str(n)),range(4)))
        self.assertTrue(all(r['total']==1200 for r in results))
        with closing(sqlite3.connect(self.path)) as con:self.assertEqual(before,list(con.iterdump()))

class AuditHttpTz2Tests(unittest.TestCase):
    def test_resource_response_is_not_http_500(self):
        fixture=audit_fixture.AuditReaderTests();fixture.setUp();self.addCleanup(fixture.doCleanups)
        with patch.object(audit_reader,'read',side_effect=audit_reader.AuditResourceError('Уточніть пошук')):
            status,body=fixture.get('?search=test')
        self.assertEqual(status,503);self.assertEqual(body['code'],'audit_resource_limit')

class ProtocolProvenanceTests(unittest.TestCase):
    def test_current_pointer_signature_and_snapshot_not_current_assignment(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)
            template=Path(__file__).parent/'templates/violation_protocols/warning.docx'
            def write(name,officer):
                from docx import Document
                document=Document(template)
                for paragraph in document.paragraphs:
                    if 'officer_name' in paragraph.text:paragraph.text='Уповноважена особа\t'+officer
                document.save(root/name)
            write('old.docx','Світлана ІСТОРИЧНА');write('new.docx','Олена НОВА')
            review=dict(generated_protocol_filename='old.docx',assigned_officer='Поточна ІНША',protocol_number='17',protocol_date='2026-10-17',generated_protocol_metadata_json=json.dumps({'protocol_subject':{'value':'Чинний погоджений короткий зміст'}}))
            before=(root/'old.docx').read_bytes()
            result=protocols.present(review,root,False)
            self.assertEqual(result['officer'],'Світлана ІСТОРИЧНА');self.assertFalse(result['json_ready'])
            self.assertEqual(result['askod'],'Чинний погоджений короткий зміст')
            review['generated_protocol_filename']='new.docx';result=protocols.present(review,root,True)
            self.assertEqual(result['officer'],'Олена НОВА');self.assertIn('new.docx',result['url']);self.assertTrue(result['json_ready'])
            self.assertEqual(before,(root/'old.docx').read_bytes())
            review['generated_protocol_filename']='missing.docx';self.assertIsNone(protocols.present(review,root,True))
            review['generated_protocol_filename']='../outside.docx';self.assertIsNone(protocols.present(review,root,True))
    def test_ambiguous_or_unknown_signature_never_uses_assignment(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            doc=Document();doc.add_paragraph('Інший історичний формат');doc.save(Path(folder)/'file.docx')
            result=protocols.present({'generated_protocol_filename':'file.docx','assigned_officer':'Поточна УО'},folder,False)
            self.assertEqual(result['officer'],'')

class CounterTz2Tests(unittest.TestCase):
    def test_all_four_conditions_and_count_list_parity_no_duplicates(self):
        fixture=counts_fixture.UiImprovementsTests();fixture.setUp();self.addCleanup(fixture.doCleanups)
        with closing(server.db()) as con:
            with con:
                con.execute('DELETE FROM violation_report_reviews');con.execute('DELETE FROM violation_reports')
                cases=[('yes','40996564','pending','2026-10-10',None,None),('foreign','other','pending','2026-10-10',None,None),('future','40996564','pending','2030-10-10',None,None),('official','40996564','satisfied','2026-10-10',None,None),('reviewed','40996564','pending','2026-10-10','reviewed',''),('legacy','40996564','pending','2026-10-10','completed',''),('stamp','40996564','pending','2026-10-10','in_review','2026-10-15')]
                for id,authority,status,end,review,completed in cases:
                    con.execute("INSERT INTO violation_reports(id,report_id,status,defendant_period_end,authority_code,raw_json,synced_at) VALUES(?,?,?,?,?,'{}','fixture')",(id,id,status,end,authority))
                    if review:con.execute("INSERT INTO violation_report_reviews(report_id,review_status,completed_at,updated_at) VALUES(?,?,?,'fixture')",(id,review,completed))
        with patch.object(server,'datetime',wraps=datetime) as clock:
            clock.now.return_value=fixture.moment
            listed=server.list_violation_reports({'response_expired':['1']})
        self.assertEqual([item['id'] for item in listed['items']],['yes'])
        self.assertEqual(listed['total'],fixture.count());self.assertEqual(listed['total'],1)

class UiTz2UnitTests(unittest.TestCase):
    def test_actual_js_queue_and_presentation(self):
        subprocess.run(['node','validation/ui_tz2_unit.cjs'],cwd=Path(__file__).parent,check=True,timeout=30)
