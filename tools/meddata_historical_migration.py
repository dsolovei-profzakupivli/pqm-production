"""Controlled MedData historical migration for WEB TEST.

Dry-run is the default.  Applying data requires the exact source hash, an
explicit writer-isolation acknowledgement and a recoverable SQLite backup.
The tool never discovers scope by date: the checked-in submission manifest is
the only authority for rows that may be changed.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sqlite3
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_SOURCE_SHA256 = "316f86077660ce53a1fc9c338ee586cc0867779b53b099db0bf4beb41b93499f"
EXPECTED_MANIFEST_COUNT = 53_803
MIGRATION_ACTOR = "migration:meddata_historical_v1"

FIELD_SPECS = (
    ("protocol_remarks", "Кваліфікаційні документи (коментарі)", "text"),
    ("protocol_number", "Номер протоколу", "text"),
    ("protocol_date", "Дата кваліфікації", "date"),
    ("publication_date", "Дата кваліфікації", "date"),
    ("protocol_officer", "Відповідальна особа (Протокол)", "officer"),
    ("review_officer", "Відповідальна особа (Публічні закупівлі)", "officer"),
    ("manager_name", "Примітки", "multiline"),
    ("contract_details", "Копія ліцензії на торгівлю лікарськими засобами", "text"),
    ("compliance_status", "Погодження комплаєнс", "compliance"),
    ("protocol_decision", "Рішення про кваліфікацію (Публічні закупівлі)", "decision"),
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized(value) -> str:
    return " ".join(str(value or "").strip().upper().split())


def officer_display(value: str) -> str:
    parts = " ".join(str(value or "").split()).split()
    return " ".join([*(part.lower().capitalize() for part in parts[:-1]), parts[-1].upper()]) if parts else ""


def text_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def multiline_value(value) -> str:
    return "\n".join(line.strip() for line in text_value(value).splitlines() if line.strip())


def date_value(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%d.%m.%Y")
    raw = text_value(value)
    for pattern in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, pattern).strftime("%d.%m.%Y")
        except ValueError:
            pass
    raise ValueError(f"Некоректна дата: {raw!r}")


def canonical_officer_map(connection: sqlite3.Connection) -> tuple[dict[str, str], set[str]]:
    rows = connection.execute("SELECT full_name FROM authorized_officers").fetchall()
    grouped: dict[str, set[str]] = {}
    for row in rows:
        value = text_value(row[0])
        grouped.setdefault(normalized(value), set()).add(value)
    ambiguous = {key for key, values in grouped.items() if len(values) != 1}
    return {key: officer_display(next(iter(values))) for key, values in grouped.items() if len(values) == 1}, ambiguous


def canonical_value(raw, kind: str, officers: dict[str, str]) -> tuple[bool, str, str]:
    source = text_value(raw)
    if not source:
        return False, "", ""
    if kind == "date":
        return True, date_value(raw), ""
    if kind == "multiline":
        return True, multiline_value(raw), ""
    if kind == "officer":
        match = officers.get(normalized(source))
        return (True, match, "") if match else (True, "", f"UNMAPPED_OFFICER:{source}")
    if kind == "compliance":
        values = {"ПОГОДЖЕНО": "approved", "НЕ ПОГОДЖЕНО": "rejected", "НЕ ВИЗНАЧЕНО": ""}
        return (True, values[normalized(source)], "") if normalized(source) in values else (True, "", f"UNKNOWN_COMPLIANCE:{source}")
    if kind == "decision":
        values = {"ТАК": "admit", "НІ": "reject", "НЕ ВИЗНАЧЕНО": ""}
        return (True, values[normalized(source)], "") if normalized(source) in values else (True, "", f"UNKNOWN_DECISION:{source}")
    return True, source, ""


def load_manifest(path: Path) -> tuple[dict, list[str]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    errors = []
    applications = payload.get("applications")
    if payload.get("version") != 1 or not isinstance(applications, dict):
        errors.append("invalid_manifest_schema")
        applications = {}
    if len(applications) != EXPECTED_MANIFEST_COUNT:
        errors.append(f"manifest_count:{len(applications)}")
    source = payload.get("source") or {}
    if source.get("sha256") != EXPECTED_SOURCE_SHA256:
        errors.append("manifest_source_hash_mismatch")
    if source.get("row_count") != EXPECTED_MANIFEST_COUNT:
        errors.append(f"manifest_source_rows:{source.get('row_count')}")
    return payload, errors


def load_source_rows(path: Path, manifest: dict) -> tuple[dict[str, dict], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    header_cells = next(sheet.iter_rows())
    headers = {text_value(cell.value): index for index, cell in enumerate(header_cells)}
    required = {source for _, source, _ in FIELD_SPECS}
    errors = [f"missing_header:{name}" for name in sorted(required - headers.keys())]
    wanted = {int(item.get("source_row") or 0): submission_id
              for submission_id, item in (manifest.get("applications") or {}).items()}
    if len(wanted) != len(manifest.get("applications") or {}):
        errors.append("duplicate_or_invalid_source_row_in_manifest")
    rows: dict[str, dict] = {}
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        submission_id = wanted.get(excel_row)
        if not submission_id:
            continue
        rows[submission_id] = {name: values[index] for name, index in headers.items()}
    missing = sorted(set(manifest.get("applications") or {}) - set(rows))
    if missing:
        errors.append(f"manifest_rows_missing_from_xlsx:{len(missing)}")
    return rows, errors


def open_readonly(path: Path) -> sqlite3.Connection:
    connection = sqlite3.connect(f"file:{path.resolve().as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA query_only=ON")
    return connection


def classify(current: str, has_source: bool, expected: str) -> str:
    if not has_source:
        return "NO_SOURCE_VALUE"
    if current == expected:
        return "SAME"
    if not current:
        return "INSERT_HISTORY"
    return "OVERWRITE_TEST_DATA"


def audit(xlsx: Path, db_path: Path, manifest_path: Path) -> tuple[dict, list[dict]]:
    source_hash = sha256(xlsx)
    manifest, errors = load_manifest(manifest_path)
    if source_hash != EXPECTED_SOURCE_SHA256:
        errors.append("source_sha256_mismatch")
    source_rows, source_errors = load_source_rows(xlsx, manifest)
    errors.extend(source_errors)
    with open_readonly(db_path) as connection:
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        if integrity != "ok":
            errors.append(f"integrity_check:{integrity}")
        officers, ambiguous_officers = canonical_officer_map(connection)
        if ambiguous_officers:
            errors.append(f"ambiguous_officer_identities:{len(ambiguous_officers)}")
        table_columns = {row[1] for row in connection.execute("PRAGMA table_info(application_fields)")}
        missing_columns = sorted({field for field, _, _ in FIELD_SPECS} - table_columns)
        if missing_columns:
            errors.append("missing_columns:" + ",".join(missing_columns))
        current_rows = {row["submission_id"]: dict(row) for row in connection.execute(
            "SELECT submission_id," + ",".join(field for field, _, _ in FIELD_SPECS) + " FROM application_fields"
        )}
        submission_ids = {row[0] for row in connection.execute("SELECT id FROM submissions")}
    manifest_ids = set(manifest.get("applications") or {})
    missing_submissions = sorted(manifest_ids - submission_ids)
    missing_fields = sorted(manifest_ids - set(current_rows))
    if missing_submissions:
        errors.append(f"unmatched_submissions:{len(missing_submissions)}")
    if missing_fields:
        errors.append(f"missing_application_fields:{len(missing_fields)}")
    plans = []
    counts = {field: Counter() for field, _, _ in FIELD_SPECS}
    holds = []
    for submission_id in sorted(manifest_ids):
        source = source_rows.get(submission_id)
        current = current_rows.get(submission_id)
        if source is None or current is None:
            continue
        for field, source_name, kind in FIELD_SPECS:
            has_source, expected, hold = canonical_value(source.get(source_name), kind, officers)
            action = "HOLD" if hold else classify(text_value(current.get(field)), has_source, expected)
            counts[field][action] += 1
            plan = {"submission_id": submission_id, "source_row": manifest["applications"][submission_id]["source_row"],
                    "field": field, "source_field": source_name, "current": text_value(current.get(field)),
                    "expected": expected, "action": action, "hold": hold}
            plans.append(plan)
            if hold:
                holds.append(plan)
    if holds:
        errors.append(f"hold:{len(holds)}")
    summary = {
        "mode": "READ_ONLY_DRY_RUN",
        "source_sha256": source_hash,
        "expected_source_sha256": EXPECTED_SOURCE_SHA256,
        "manifest_count": len(manifest_ids),
        "expected_manifest_count": EXPECTED_MANIFEST_COUNT,
        "resolved": len(manifest_ids) - len(missing_submissions),
        "ambiguous": 0,
        "unmatched": len(missing_submissions),
        "hold": len(holds),
        "integrity_check": integrity,
        "field_actions": {field: dict(sorted(counter.items())) for field, counter in counts.items()},
        "hard_stop_reasons": sorted(set(errors)),
        "safe_to_apply": not errors,
        "scope_authority": "explicit_submission_manifest_only",
    }
    return summary, plans


def backup_database(db_path: Path, backup_dir: Path) -> dict:
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    destination = backup_dir / f"pqm_before_meddata_{stamp}.sqlite3"
    source = sqlite3.connect(db_path)
    target = sqlite3.connect(destination)
    try:
        source.backup(target)
    finally:
        target.close(); source.close()
    return {"path": str(destination), "sha256": sha256(destination),
            "size": destination.stat().st_size,
            "wal_present": db_path.with_name(db_path.name + "-wal").exists(),
            "shm_present": db_path.with_name(db_path.name + "-shm").exists()}


def apply(db_path: Path, plans: list[dict], backup_dir: Path) -> dict:
    backup = backup_database(db_path, backup_dir)
    connection = sqlite3.connect(db_path, timeout=60)
    changed = Counter()
    stamp = datetime.now(timezone.utc).isoformat()
    try:
        connection.execute("PRAGMA foreign_keys=ON")
        connection.execute("BEGIN IMMEDIATE")
        for row in plans:
            if row["action"] not in {"INSERT_HISTORY", "OVERWRITE_TEST_DATA"}:
                continue
            cursor = connection.execute(
                f"UPDATE application_fields SET {row['field']}=? WHERE submission_id=? AND COALESCE({row['field']},'')=?",
                (row["expected"], row["submission_id"], row["current"]),
            )
            if cursor.rowcount != 1:
                raise RuntimeError(f"Concurrent change detected: {row['submission_id']} {row['field']}")
            connection.execute("""INSERT INTO audit_log
              (submission_id,changed_at,changed_by,field_name,old_value,new_value)
              VALUES (?,?,?,?,?,?)""", (row["submission_id"], stamp, MIGRATION_ACTOR,
              row["field"], row["current"], row["expected"]))
            changed[row["field"]] += 1
        connection.commit()
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    if integrity != "ok":
        raise RuntimeError(f"Post-apply integrity_check failed: {integrity}")
    return {"backup": backup, "changed": dict(changed), "integrity_check": integrity}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--db", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--confirm-source-sha", default="")
    parser.add_argument("--writers-stopped", action="store_true")
    parser.add_argument("--backup-dir", type=Path)
    args = parser.parse_args()
    summary, plans = audit(args.xlsx, args.db, args.manifest)
    result = {"summary": summary}
    if args.apply:
        if args.confirm_source_sha.lower() != EXPECTED_SOURCE_SHA256:
            raise SystemExit("APPLY BLOCKED: exact --confirm-source-sha is required")
        if not args.writers_stopped:
            raise SystemExit("APPLY BLOCKED: stop WEB writers and pass --writers-stopped")
        if not args.backup_dir:
            raise SystemExit("APPLY BLOCKED: --backup-dir is required")
        if not summary["safe_to_apply"]:
            raise SystemExit("APPLY BLOCKED: " + ", ".join(summary["hard_stop_reasons"]))
        result["apply"] = apply(args.db, plans, args.backup_dir)
        result["summary"]["mode"] = "APPLY"
    raw = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(raw, encoding="utf-8")
    print(raw, end="")


if __name__ == "__main__":
    main()
