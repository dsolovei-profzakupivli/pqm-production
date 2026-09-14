"""Build the explicit MedData application manifest without modifying SQLite/XLSX."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def clean(value) -> str:
    return " ".join(str(value or "").strip().casefold().split())


def iso_date(value) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    for pattern in ("%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise ValueError(f"Некоректна дата MedData: {text!r}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--db", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    source_hash = hashlib.sha256(args.xlsx.read_bytes()).hexdigest()
    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    sheet = workbook.active
    headers = {str(cell.value or "").strip(): index for index, cell in enumerate(next(sheet.iter_rows()))}
    required = {
        "Код ЄДРПОУ", "CPV код", "Назва фреймворку", "Дата розміщення документів",
        "Номер протоколу", "Дата кваліфікації", "Рішення про кваліфікацію (Публічні закупівлі)",
    }
    missing = sorted(required - headers.keys())
    if missing:
        raise RuntimeError("В XLSX відсутні колонки: " + ", ".join(missing))

    connection = sqlite3.connect(f"file:{args.db.resolve().as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    candidates = defaultdict(list)
    for row in connection.execute("""SELECT s.id,s.supplier_code,SUBSTR(s.date_published,1,10) placement_date,
      COALESCE(f.dk_code,'') cpv,COALESCE(f.title,'') framework_title,
      COALESCE(af.protocol_number,'') protocol_number,COALESCE(af.protocol_date,'') protocol_date,
      COALESCE(q.status,'pending') qualification_status
      FROM submissions s LEFT JOIN frameworks f ON f.id=s.framework_id
      LEFT JOIN application_fields af ON af.submission_id=s.id
      LEFT JOIN qualifications q ON q.id=s.qualification_id"""):
        key = (re.sub(r"\D", "", row["supplier_code"] or ""), row["placement_date"], row["cpv"])
        candidates[key].append(dict(row))

    applications = {}
    unresolved = []
    decision_status = {"так": "active", "ні": "unsuccessful", "": "pending"}
    source_rows = 0
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        source_rows += 1
        code = re.sub(r"\D", "", str(values[headers["Код ЄДРПОУ"]] or ""))
        cpv = str(values[headers["CPV код"]] or "").strip()
        placement = iso_date(values[headers["Дата розміщення документів"]])
        title = clean(values[headers["Назва фреймворку"]])
        found = candidates.get((code, placement, cpv), [])
        title_matches = [item for item in found if clean(item["framework_title"]) == title]
        found = title_matches or found
        if len(found) > 1:
            number = str(values[headers["Номер протоколу"]] or "").strip()
            raw_protocol_date = values[headers["Дата кваліфікації"]]
            protocol_date = iso_date(raw_protocol_date) if raw_protocol_date else ""
            status = decision_status.get(clean(values[headers["Рішення про кваліфікацію (Публічні закупівлі)"]]))
            narrowed = [item for item in found
                        if str(item["protocol_number"] or "").strip() == number
                        and (not protocol_date or iso_date(item["protocol_date"]) == protocol_date)
                        and (status is None or item["qualification_status"] == status)]
            found = narrowed
        if len(found) != 1:
            unresolved.append({"source_row": excel_row, "supplier_code": code,
                               "placement_date": placement, "cpv": cpv,
                               "candidate_count": len(found)})
            continue
        submission_id = found[0]["id"]
        if submission_id in applications:
            raise RuntimeError(f"Два source rows зіставлено з {submission_id}")
        applications[submission_id] = {"source_system": "MedData", "source_row": excel_row}

    if unresolved:
        raise RuntimeError("Неоднозначне зіставлення: " + json.dumps(unresolved, ensure_ascii=False))
    payload = {
        "version": 1,
        "source": {"filename": args.xlsx.name, "sha256": source_hash,
                   "sheet": sheet.title, "row_count": source_rows},
        "applications": dict(sorted(applications.items())),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({"source_rows": source_rows, "matched": len(applications),
                      "sha256": source_hash, "output": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
