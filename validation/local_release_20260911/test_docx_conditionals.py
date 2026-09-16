import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt
from lxml import etree
import schema_catalog
import template_catalog
from docx_conditionals import render
from protocol_template import NS, paragraph_text
from template_conditions import ConditionalError

TYPE='nazk_supplier_request'
AMCU_TYPE='amcu_exclusion_protocol'
FOP='{{#if supplier.entity_type == "individual_entrepreneur"}}'
LEGAL='{{#if supplier.entity_type == "legal_entity"}}'
END='{{/if}}'
AMCU_SINGLE='{{#if amcu.is_single_decision == "true"}}'
AMCU_MULTIPLE='{{#if amcu.has_multiple_decisions == "true"}}'


class ConditionalTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fields=template_catalog.validate(template_catalog.load(),schema_catalog.catalog(Path('data/pqm.sqlite3')))

    def setUp(self):
        self.folder=tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.source=Path(self.folder.name)/'source.docx'
        self.output=Path(self.folder.name)/'result.docx'
        self.context={'supplier.entity_type':'legal_entity','supplier.email':'test@example.com','supplier.name':'Назва'}

    def fixture(self, paragraphs):
        d=Document()
        d.styles['Normal'].font.name='Times New Roman'
        d.styles['Normal'].font.size=Pt(12)
        for text in paragraphs:d.add_paragraph(text)
        d.save(self.source)
        return d

    def build(self):
        return render(self.source,self.output,self.context,self.fields,TYPE)

    def test_both_branches_scalars_no_markers_or_empty_paragraphs(self):
        self.fixture(['Назва: {{supplier.name}}',FOP,'Для ФОП',END,LEGAL,'Для ЮО',END,'Кінець'])
        before=self.source.read_bytes()
        for entity,expected,absent in [('individual_entrepreneur','Для ФОП','Для ЮО'),('legal_entity','Для ЮО','Для ФОП')]:
            self.context['supplier.entity_type']=entity
            self.build()
            self.assertEqual([p.text for p in Document(self.output).paragraphs],['Назва: Назва',expected,'Кінець'])
            self.assertEqual(self.source.read_bytes(),before)
            self.assertNotIn(absent,'\n'.join(p.text for p in Document(self.output).paragraphs))

    def test_invalid_context_no_output(self):
        self.fixture([FOP,'FOP',END])
        for value in (None,'unknown',1):
            self.context['supplier.entity_type']=value
            with self.assertRaises(ConditionalError):self.build()
            self.assertFalse(self.output.exists())
        del self.context['supplier.entity_type']
        with self.assertRaises(ConditionalError):self.build()

    def test_bad_syntax_no_output(self):
        for paragraphs in ([FOP,'x'],[END],[FOP,LEGAL,'x',END,END],
            ['{{#if supplier.entity_type != "legal_entity"}}','x',END],
            ['{{#if supplier.entity_type == "legal_entity" or true}}','x',END],
            [FOP,'{{else}}',END],['inline '+FOP,'x',END],
            ['{{#if supplier.entity_type == "unknown"}}','x',END],
            ['{{#if missing.field == "x"}}','x',END]):
            with self.subTest(paragraphs=paragraphs):
                self.fixture(paragraphs)
                with self.assertRaises(ConditionalError):self.build()
                self.assertFalse(self.output.exists())

    def test_unavailable_and_non_string_field(self):
        for marker in ('{{#if warning.references == "x"}}','{{#if system.today == "x"}}'):
            self.fixture([marker,'x',END])
            with self.assertRaises(ConditionalError):self.build()

    def test_split_markers_and_scalar_style_preserved(self):
        d=Document()
        p=d.add_paragraph();p.add_run(FOP[:10]);p.add_run(FOP[10:])
        p=d.add_paragraph();p.add_run('{{supplier.').bold=True;p.add_run('name}}')
        d.add_paragraph(END);d.add_paragraph('Tail');d.save(self.source)
        self.context['supplier.entity_type']='individual_entrepreneur'
        self.build();p=Document(self.output).paragraphs[0]
        self.assertEqual(p.text,'Назва');self.assertTrue(p.runs[0].bold)

    def test_whole_table_removed_headers_supported(self):
        d=Document();d.add_paragraph(FOP)
        d.add_table(rows=1,cols=1).cell(0,0).text='FOP table'
        d.add_paragraph(END);d.add_paragraph('Tail')
        h=d.sections[0].header;h.paragraphs[0].text=LEGAL;h.add_paragraph('Header');h.add_paragraph(END)
        d.save(self.source);self.build();out=Document(self.output)
        self.assertEqual(len(out.tables),0)
        self.assertEqual([p.text for p in out.paragraphs],['Tail'])
        self.assertEqual([p.text for p in out.sections[0].header.paragraphs],['Header'])

    def test_cross_cell_markers_rejected(self):
        d=Document();table=d.add_table(rows=1,cols=2)
        table.cell(0,0).text=FOP;table.cell(0,1).text=END;d.save(self.source)
        with self.assertRaises(ConditionalError):self.build()

    def test_required_email_and_atomic_failure(self):
        self.fixture([LEGAL,'x',END]);self.build();before=self.output.read_bytes()
        self.context['supplier.email']=None
        with self.assertRaises(ConditionalError):self.build()
        self.assertEqual(self.output.read_bytes(),before)

    def test_generic_string_condition_not_nazk_specific(self):
        self.fixture(['{{#if supplier.name == "Назва"}}','Matched',END,
                      '{{#if supplier.name == "Other"}}','','Not matched',END])
        self.build()
        self.assertEqual([p.text for p in Document(self.output).paragraphs],['Matched'])

    def test_false_branch_still_validated(self):
        self.fixture([FOP,'{{unknown.field}}',END])
        with self.assertRaises(ConditionalError):self.build()
        self.assertFalse(self.output.exists())
        from docx_conditionals import condition_keys, retained_scalar_keys
        self.fixture([FOP,'{{manager.full_name_genitive}}',END,LEGAL,'{{supplier.name_genitive}}',END])
        self.assertEqual(condition_keys(self.source,self.fields,TYPE),{'supplier.entity_type'})
        context={'supplier.entity_type':'individual_entrepreneur'}
        self.assertEqual(retained_scalar_keys(self.source,self.fields,TYPE,context),{'manager.full_name_genitive'})

    def test_same_cell_blocks_keep_other_content(self):
        d=Document();cell=d.add_table(rows=1,cols=1).cell(0,0)
        cell.paragraphs[0].text=FOP;cell.add_paragraph('Excluded');cell.add_paragraph(END);cell.add_paragraph('Kept')
        d.save(self.source);self.build()
        self.assertEqual(Document(self.output).tables[0].cell(0,0).text,'Kept')

    def test_catalog_scanner_conditional_gate(self):
        self.fixture([LEGAL,'{{supplier.name}}',END])
        report=template_catalog.scan_docx(self.source,self.fields,TYPE)
        self.assertTrue(report['can_activate_canonical'])
        self.fixture([LEGAL,'x'])
        report=template_catalog.scan_docx(self.source,self.fields,TYPE)
        self.assertFalse(report['can_activate_canonical']);self.assertTrue(report['conditional_errors'])

    def test_existing_four_packages_unchanged(self):
        from violation_protocol_docx import TEMPLATES
        for key,path in TEMPLATES.items():
            before=path.read_bytes()
            report=template_catalog.scan_docx(path,self.fields,template_catalog.RUNTIME_TYPES[key],key)
            self.assertFalse(report['conditional_errors'])
            self.assertFalse(report['unknown'])
            self.assertEqual(before,path.read_bytes())

    def test_amcu_repeat_paragraph_group_and_date_format(self):
        self.fixture(['Протокол № {{decision.number}} від {{decision.date}}','До блоку','{{#repeat amcu.decisions[]}}',
                      'Рішення № {{number}} від {{date}} · {{authority}}',
                      'Витяг: {{extract_url}}','{{/repeat}}','Після блоку'])
        context={'decision.number':'701','decision.date':'2026-09-14','amcu.decisions[]':[
          {'number':'72/130-р/к','date':'2026-09-11','authority':'АМКУ','extract_url':'https://one'},
          {'number':'73/130-р/к','date':'2026-09-12','authority':'Комісія','extract_url':'https://two'}]}
        render(self.source,self.output,context,self.fields,AMCU_TYPE)
        self.assertEqual([p.text for p in Document(self.output).paragraphs],[
          'Протокол № 701 від 14.09.2026','До блоку','Рішення № 72/130-р/к від 11.09.2026 · АМКУ','Витяг: https://one',
          'Рішення № 73/130-р/к від 12.09.2026 · Комісія','Витяг: https://two','Після блоку'])
        report=template_catalog.scan_docx(self.source,self.fields,AMCU_TYPE)
        self.assertTrue(report['can_activate_canonical'],report)

    def test_amcu_repeat_url_becomes_external_clickable_hyperlink(self):
        self.fixture(['{{#repeat amcu.decisions[]}}','Посилання на витяг: {{extract_url}}','{{/repeat}}'])
        render(self.source,self.output,{'amcu.decisions[]':[
          {'number':'1','date':'2026-09-11','authority':'АМКУ','extract_url':'https://example.test/one'},
          {'number':'2','date':'2026-09-12','authority':'АМКУ','extract_url':'https://example.test/two'}]},
          self.fields,AMCU_TYPE)
        with ZipFile(self.output) as package:
            root=etree.fromstring(package.read('word/document.xml'))
            links=root.xpath('.//w:hyperlink',namespaces=NS)
            self.assertEqual(['https://example.test/one','https://example.test/two'],[
                paragraph_text(link) for link in links])
            relationships=etree.fromstring(package.read('word/_rels/document.xml.rels'))
            by_id={item.get('Id'):item for item in relationships}
            self.assertEqual(['https://example.test/one','https://example.test/two'],[
                by_id[link.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')].get('Target')
                for link in links])
            self.assertTrue(all(by_id[link.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')].get('TargetMode')=='External'
                                for link in links))
        self.assertEqual([p.text for p in Document(self.output).paragraphs],[
            'Посилання на витяг: https://example.test/one','Посилання на витяг: https://example.test/two'])

    def test_amcu_repeat_linked_reference_hides_raw_url_and_links_each_item(self):
        self.fixture(['{{#repeat amcu.decisions[]}}','{{linked_reference}}','{{/repeat}}'])
        render(self.source,self.output,{'amcu.decisions[]':[
          {'number':'1','date':'2026-09-11','authority':'','extract_url':'https://example.test/one','linked_reference':'від 11.09.2026 № 1'},
          {'number':'2','date':'2026-09-12','authority':'','extract_url':'https://example.test/two','linked_reference':'від 12.09.2026 № 2'}]},
          self.fields,AMCU_TYPE)
        with ZipFile(self.output) as package:
            root=etree.fromstring(package.read('word/document.xml'))
            self.assertEqual([paragraph_text(link) for link in root.xpath('.//w:hyperlink',namespaces=NS)],
                             ['від 11.09.2026 № 1','від 12.09.2026 № 2'])
            text='\n'.join(paragraph_text(p) for p in root.xpath('.//w:p',namespaces=NS))
            self.assertNotIn('https://',text);self.assertNotIn('посилання на витяг',text.casefold())
            rels=etree.fromstring(package.read('word/_rels/document.xml.rels'))
            by_id={row.get('Id'):row.get('Target') for row in rels}
            self.assertEqual([by_id[link.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
                              for link in root.xpath('.//w:hyperlink',namespaces=NS)],
                             ['https://example.test/one','https://example.test/two'])
            for link in root.xpath('.//w:hyperlink',namespaces=NS):
                properties=link.xpath('./w:r/w:rPr',namespaces=NS)[0]
                self.assertEqual(properties.xpath('./w:color/@w:val',namespaces=NS),['0563C1'])
                self.assertEqual(properties.xpath('./w:u/@w:val',namespaces=NS),['single'])

    def test_amcu_singular_plural_conditions_keep_repeat_hyperlinks_and_no_markers(self):
        self.fixture([
          AMCU_SINGLE,'встановлено наявність такого рішення:',END,
          AMCU_MULTIPLE,'встановлено наявність таких рішень:',END,
          '{{#repeat amcu.decisions[]}}','{{linked_reference}}','{{/repeat}}',
          AMCU_SINGLE,'На підставі зазначеного рішення постачальника.',END,
          AMCU_MULTIPLE,'На підставі зазначених рішень постачальника.',END,
          AMCU_SINGLE,'Штраф накладено на підставі такого рішення:',END,
          AMCU_MULTIPLE,'Штраф накладено на підставі таких рішень:',END])
        for count in (1,2,3):
            with self.subTest(count=count):
                items=[{'number':str(index),'date':f'2026-09-{10+index:02d}','authority':'',
                        'extract_url':f'https://example.test/{index}',
                        'linked_reference':f'від {10+index:02d}.09.2026 № {index}'}
                       for index in range(1,count+1)]
                context={'amcu.is_single_decision':'true' if count==1 else 'false',
                         'amcu.has_multiple_decisions':'true' if count>1 else 'false',
                         'amcu.decisions[]':items}
                render(self.source,self.output,context,self.fields,AMCU_TYPE)
                with ZipFile(self.output) as package:
                    root=etree.fromstring(package.read('word/document.xml'))
                    text='\n'.join(paragraph_text(p) for p in root.xpath('.//w:p',namespaces=NS))
                    self.assertNotIn('{{',text);self.assertNotIn('}}',text)
                    if count==1:
                        self.assertEqual(text.count('такого рішення'),2)
                        self.assertIn('зазначеного рішення',text)
                        self.assertNotIn('таких рішень',text);self.assertNotIn('зазначених рішень',text)
                    else:
                        self.assertEqual(text.count('таких рішень'),2)
                        self.assertIn('зазначених рішень',text)
                        self.assertNotIn('такого рішення',text);self.assertNotIn('зазначеного рішення',text)
                    links=root.xpath('.//w:hyperlink',namespaces=NS)
                    self.assertEqual(len(links),count)
                    relationships=etree.fromstring(package.read('word/_rels/document.xml.rels'))
                    by_id={row.get('Id'):row for row in relationships}
                    self.assertEqual(
                      [by_id[link.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')].get('Target')
                       for link in links],
                      [item['extract_url'] for item in items])
                    self.assertTrue(all(by_id[link.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')].get('TargetMode')=='External'
                                        for link in links))

    def test_empty_amcu_repeat_removes_complete_group(self):
        self.fixture(['До','{{#repeat amcu.decisions[]}}','№ {{number}}','{{/repeat}}','Після'])
        render(self.source,self.output,{'amcu.decisions[]':[]},self.fields,AMCU_TYPE)
        self.assertEqual([p.text for p in Document(self.output).paragraphs],['До','Після'])

    def test_repeat_scope_is_strict_and_atomic(self):
        for paragraphs in (["{{#repeat amcu.decisions[]}}",'{{unknown}}','{{/repeat}}'],
                           ['{{number}}'],['{{#repeat amcu.decisions[]}}','{{number}}']):
            with self.subTest(paragraphs=paragraphs):
                self.fixture(paragraphs)
                with self.assertRaises(ConditionalError):
                    render(self.source,self.output,{'amcu.decisions[]':[]},self.fields,AMCU_TYPE)
                self.assertFalse(self.output.exists())
