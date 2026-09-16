import hashlib
import os
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import violation_protocol_docx as generator


EXPECTED_HASHES = {
    "warning": "33d1a471142377ea01ac11b403708d77718e506e814103215d66a8954d58e32e",
    # Reviewed NEXT packaged asset only; the operator's runtime copy is preserved.
    "decline_p49_1_2": "525cb13508d5db3c994fde50b342b79fdc1e5b3673f4fa1e09155efb87e81016",
    "decline_p49_3": "5bf29f6a84f7136a60a48f16b9b0f8722451071bdf1654446b79e61805234622",
}

MINISTRY_URLS = (
    "https://www.me.gov.ua/InfoRez/Details?id=3d8b5293-1542-45e7-8cab-60768b9ecc09&lang=uk-UA",
    "https://me.gov.ua/InfoRez/Details?id=1c50d66b-a34f-4b83-8ae3-e1fdea208d80&lang=uk-UA",
    "https://me.gov.ua/InfoRez/Details?id=011d5df6-768e-46e9-9f66-86a71737584d&lang=uk-UA",
)

BASE_VALUES = {
    "protocol_number": "TEST-1", "protocol_date": "08.09.2026", "report_id": "UA-D-TEST",
    "procurement_id": "UA-2026-TEST", "procurement_date": "01.09.2026", "cpv_category": "44110000-4",
    "customer_name": "Замовник", "customer_name_genitive": "Замовника",
    "customer_name_accusative": "Замовника", "customer_code": "12345678",
    "supplier_name": "Постачальник", "supplier_short_name": "Постачальник",
    "supplier_name_genitive": "Постачальника", "supplier_name_dative": "Постачальнику",
    "supplier_name_accusative": "Постачальника", "supplier_code": "87654321",
    "supplier_code_label": "код ЄДРПОУ", "officer_name": "Тестова УО",
    "p49_reference": "пп. 1 п. 49", "reason_label": "Підстава", "reason_text": "Опис підстави",
    "violation_description": "Фактичний опис порушення", "winner_date": "01.09.2026",
    "supplier_deadline": "05.09.2026", "rejection_date": "06.09.2026",
    "rejection_reason": "Фактична підстава відхилення", "refusal_date": "04.09.2026",
    "refusal_outgoing_number": "42", "refusal_document": "https://example.test/refusal",
    "supplier_response": "Фактичне пояснення постачальника", "contract_date": "03.09.2026",
    "contract_number": "C-77",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def all_text(path: Path) -> str:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(chunks)


class ViolationProtocolDocxTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)

    def tearDown(self):
        self.temp.cleanup()

    def build(self, protocol_type="warning", customer=None, supplier=None):
        output = Path(self.temp.name) / f"{protocol_type}.docx"
        generator.build_violation_protocol_docx(
            protocol_type, output,
            dict(BASE_VALUES),
            "ОСТАТОЧНЕ ОБҐРУНТУВАННЯ УО ДОСЛІВНО",
            customer or [], supplier or [],
            {"has_written_refusal": True, "has_contract": True, "has_contract_security": True,
             "has_civil_code_basis": True, "has_court_decision": True,
             "has_supplier_response": True, "has_supplier_documents": bool(supplier),
             "has_customer_documents": bool(customer)},
        )
        return output

    def test_approved_source_templates_are_unchanged(self):
        for key, expected in EXPECTED_HASHES.items():
            self.assertEqual(sha256(generator.PACKAGED_TEMPLATE_DIR / f"{key}.docx"), expected)

    def test_all_protocol_types_generate_without_placeholders(self):
        for protocol_type in EXPECTED_HASHES:
            output = self.build(protocol_type)
            text = all_text(output)
            self.assertNotIn("{{", text)
            self.assertIn("ОСТАТОЧНЕ ОБҐРУНТУВАННЯ УО ДОСЛІВНО", text)

    def test_replacement_does_not_add_highlight_and_keeps_static_template_highlight(self):
        document = Document()
        paragraph = document.add_paragraph()
        static = paragraph.add_run("Навмисно виділено: ")
        static.font.highlight_color = WD_COLOR_INDEX.TURQUOISE
        placeholder = paragraph.add_run("{{ supplier_name_genitive }}")
        self.assertIsNone(placeholder.font.highlight_color)

        generator._replace_in_paragraph(
            paragraph, {"supplier_name_genitive": "КОНТРОЛЬНОГО ПОСТАЧАЛЬНИКА"})

        static_result = next(run for run in paragraph.runs if "Навмисно виділено" in run.text)
        resolved = next(run for run in paragraph.runs if "КОНТРОЛЬНОГО ПОСТАЧАЛЬНИКА" in run.text)
        self.assertEqual(static_result.font.highlight_color, WD_COLOR_INDEX.TURQUOISE)
        self.assertIsNone(resolved.font.highlight_color)

    def test_p49_reference_and_adjacent_separator_have_no_service_shading(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("передбаченого ")
        placeholder = paragraph.add_run("{{ p49_reference }}")
        shaded_space = paragraph.add_run(" ")
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), "FF9900")
        shaded_space._r.get_or_add_rPr().append(shading)
        paragraph.add_run("Порядку № 822")
        self.assertIsNone(placeholder._r.rPr)

        generator._replace_in_paragraph(paragraph, {"p49_reference": "пп. 1 п. 49"})

        resolved = next(run for run in paragraph.runs if "пп. 1 п. 49" in run.text)
        self.assertIsNone(resolved.font.highlight_color)
        self.assertTrue(resolved._r.rPr is None or resolved._r.rPr.find(qn("w:shd")) is None)
        static_space = next(run for run in paragraph.runs if run.text == " ")
        self.assertTrue(
            static_space._r.rPr is None or static_space._r.rPr.find(qn("w:shd")) is None
        )
        self.assertEqual(paragraph.text, "передбаченого пп. 1 п. 49 Порядку № 822")

    def test_generated_declined_names_have_no_service_highlight(self):
        values = {
            **BASE_VALUES,
            "customer_name_genitive": "КОНТРОЛЬНОГО ЗАМОВНИКА РОДОВИЙ",
            "customer_name_accusative": "КОНТРОЛЬНОГО ЗАМОВНИКА ЗНАХІДНИЙ",
            "supplier_name_genitive": "КОНТРОЛЬНОГО ПОСТАЧАЛЬНИКА РОДОВИЙ",
            "supplier_name_dative": "КОНТРОЛЬНОМУ ПОСТАЧАЛЬНИКУ ДАВАЛЬНИЙ",
            "supplier_name_accusative": "КОНТРОЛЬНОГО ПОСТАЧАЛЬНИКА ЗНАХІДНИЙ",
        }
        output = Path(self.temp.name) / "no-service-highlight.docx"
        generator.build_violation_protocol_docx(
            "warning", output, values, "ОСТАТОЧНЕ ОБҐРУНТУВАННЯ УО ДОСЛІВНО", [], [],
            {"has_written_refusal": True, "has_contract": True, "has_contract_security": True,
             "has_civil_code_basis": True, "has_court_decision": True,
             "has_supplier_response": True},
        )
        document = Document(output)
        expected = {values[key] for key in (
            "customer_name_genitive", "customer_name_accusative", "supplier_name_genitive",
            "supplier_name_dative", "supplier_name_accusative")}
        found = set()
        for paragraph in generator._all_paragraphs(document):
            for run in paragraph.runs:
                for value in expected:
                    if value in run.text:
                        found.add(value)
                        self.assertNotEqual(run.font.highlight_color, WD_COLOR_INDEX.YELLOW, value)
        self.assertEqual(found, expected)

    def test_no_canonical_placeholder_value_gets_automatic_highlight_or_shading(self):
        for protocol_type in EXPECTED_HASHES:
            output = self.build(protocol_type)
            document = Document(output)
            expected = {generator._presentation_text(value).strip() for value in BASE_VALUES.values()
                        if len(str(value).strip()) >= 4}
            for paragraph in generator._all_paragraphs(document):
                for run in paragraph.runs:
                    if not any(value in run.text for value in expected):
                        continue
                    self.assertIsNone(run.font.highlight_color, (protocol_type, run.text))
                    self.assertTrue(run._r.rPr is None or run._r.rPr.find(qn("w:shd")) is None,
                                    (protocol_type, run.text))

    def test_reason_block_has_three_distinct_semantic_parts_and_template_styles(self):
        values = {**BASE_VALUES,
                  "reason_label": "Підпункт 1 пункту 49 Постанови Кабінету Міністрів України",
                  "reason_text": "Офіційний нормативний опис Prozorro",
                  "violation_description": "Оригінальне пояснення Замовника"}
        for protocol_type in EXPECTED_HASHES:
            output = Path(self.temp.name) / f"reason-{protocol_type}.docx"
            generator.build_violation_protocol_docx(
                protocol_type, output, values, "Перше речення.\n\nДруге речення.", [], [],
                {"has_written_refusal": True, "has_contract": True,
                 "has_supplier_response": True})
            document = Document(output)
            reason_table = document.tables[1]
            label_cell, value_cell = reason_table.rows[0].cells[0], reason_table.rows[0].cells[-1]
            description_cell = reason_table.rows[1].cells[0]
            self.assertEqual(label_cell.text, "Причина звернення:")
            self.assertEqual([p.text for p in value_cell.paragraphs], [
                values["reason_label"], values["reason_text"]])
            self.assertEqual(description_cell.text, "«Оригінальне пояснення Замовника»")
            for run in label_cell.paragraphs[0].runs + value_cell.paragraphs[0].runs:
                self.assertEqual(run.font.name, "Times New Roman")
                self.assertEqual(run.font.size.pt, 10)
                self.assertTrue(run.bold)
            for run in value_cell.paragraphs[1].runs:
                self.assertEqual(run.font.name, "Times New Roman")
                self.assertEqual(run.font.size.pt, 10)
                self.assertFalse(bool(run.bold))
            for run in description_cell.paragraphs[0].runs:
                self.assertEqual(run.font.name, "Times New Roman")
                self.assertEqual(run.font.size.pt, 10)
                self.assertTrue(run.italic)

    def test_justification_preserves_paragraphs_and_has_explicit_effective_formatting(self):
        justification = ("\tВідповідно до пп. 2 п. 49 Порядку № 822 застосовується правило.\r\n\r\n"
                         "  Гранична дата 05.09.2026. https://example.test/source\r"
                         "Наступний рядок.\n\n")
        output = Path(self.temp.name) / "justification.docx"
        generator.build_violation_protocol_docx(
            "decline_p49_1_2", output, BASE_VALUES, justification, [], [],
            {"has_written_refusal": True, "has_contract": True,
             "has_supplier_response": True})
        document = Document(output)
        row = next(row for table in document.tables for row in table.rows
                   if "Обґрунтування рішення" in " ".join(cell.text for cell in row.cells))
        label, target = row.cells[0], row.cells[-1]
        self.assertEqual(len(target.paragraphs), 3)
        expected_justification = generator._presentation_text(
            "Відповідно до пп. 2 п. 49 Порядку № 822 застосовується правило.\n"
            "Гранична дата 05.09.2026. https://example.test/source\n"
            "Наступний рядок.")
        self.assertEqual("\n".join(p.text for p in target.paragraphs), expected_justification)
        self.assertTrue(all(p.text for p in target.paragraphs))
        for paragraph in target.paragraphs:
            self.assertEqual(paragraph.alignment, generator.WD_ALIGN_PARAGRAPH.JUSTIFY)
            self.assertAlmostEqual(paragraph.paragraph_format.first_line_indent.cm, 1.0, places=2)
            self.assertEqual(paragraph.paragraph_format.space_before.pt, 0)
            self.assertEqual(paragraph.paragraph_format.space_after.pt, 0)
        self.assertTrue(all(run.bold for run in label.paragraphs[0].runs if run.text))
        self.assertTrue(any(run.bold and "пп.\u00a02 п.\u00a049" in run.text
                            for run in target.paragraphs[0].runs))
        # Every generated paragraph inherits the complete source paragraph pPr;
        # formatting remains controlled by the approved DOCX template.
        ppr_xml = [p._p.pPr.xml if p._p.pPr is not None else "" for p in target.paragraphs]
        self.assertTrue(all(xml == ppr_xml[0] for xml in ppr_xml[1:]))
        with zipfile.ZipFile(output) as package:
            xml = package.read("word/document.xml")
        from lxml import etree
        root = etree.fromstring(xml)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        justification_row = next(tr for tr in root.xpath(".//w:tr", namespaces=ns)
                                 if "Обґрунтування рішення" in "".join(tr.itertext()))
        right_cell = justification_row.xpath("./w:tc[last()]", namespaces=ns)[0]
        for run in right_cell.xpath(".//w:r", namespaces=ns):
            text = "".join(run.xpath(".//w:t/text()", namespaces=ns))
            if not text:
                continue
            fonts = run.find("w:rPr/w:rFonts", ns)
            size = run.find("w:rPr/w:sz", ns)
            self.assertIsNotNone(fonts)
            self.assertEqual(fonts.get(qn("w:ascii")), "Times New Roman")
            self.assertIsNotNone(size)
            self.assertEqual(size.get(qn("w:val")), "24")

    def test_local_unavailable_flag_does_not_claim_source_file_is_damaged(self):
        output = self.build(customer=[{
            "id": "c1", "title": "Ламаний файл.docx", "url": "https://example.test/broken",
            "file_unavailable": True,
        }])
        text = all_text(output)
        self.assertNotIn("технічно пошкоджен", text.lower())
        self.assertNotIn("0 байт", text.lower())
        with zipfile.ZipFile(output) as package:
            relationships = package.read("word/_rels/document.xml.rels").decode("utf-8")
        self.assertIn("https://example.test/broken", relationships)

    def test_supplier_unavailable_adds_no_generated_wording(self):
        output = self.build(supplier=[{
            "id": "s1", "title": "Пояснення.pdf", "url": "https://example.test/supplier",
            "file_unavailable": True,
        }])
        text = all_text(output)
        self.assertNotIn("файл не відкривається", text.lower())
        self.assertNotIn("технічно пошкоджен", text.lower())
        with zipfile.ZipFile(output) as package:
            relationships = package.read("word/_rels/document.xml.rels").decode("utf-8")
        self.assertIn("https://example.test/supplier", relationships)

    def test_optional_contract_is_removed_but_supplier_outcome_rows_remain(self):
        output = Path(self.temp.name) / "no_optional.docx"
        generator.build_violation_protocol_docx(
            "decline_p49_1_2", output, {**BASE_VALUES, "supplier_response": "",
            "contract_date": "", "contract_number": ""},
            "Збережене обґрунтування", [], [],
            {"has_written_refusal": False, "has_contract": False,
             "has_supplier_response": False, "has_supplier_documents": False,
             "has_customer_documents": False},
        )
        text = all_text(output)
        self.assertNotIn("C-77", text)
        self.assertNotIn("Фактичне пояснення постачальника", text)
        self.assertIn("Відповідь постачальника на звернення:", text)
        self.assertIn("Документи подані постачальником:", text)
        document = Document(output)
        outcomes = []
        for table in document.tables:
            for row in table.rows:
                if any(label in " ".join(cell.text for cell in row.cells)
                       for label in ("Відповідь постачальника", "Документи подані постачальником")):
                    cells = []
                    seen = set()
                    for cell in row.cells:
                        if cell._tc not in seen:
                            seen.add(cell._tc); cells.append(cell)
                    outcomes.append(cells[-1].text.strip())
        self.assertEqual(outcomes, ["не надано", "не надано"])
        self.assertNotIn("{{", text)

    def test_supplier_response_and_documents_have_four_required_states_in_all_templates(self):
        combinations = (
            ("", [], 2),
            ("Повний фактичний текст відповіді", [], 1),
            ("", [{"id": "s1", "title": "доказ.pdf", "url": "https://example.test/s1"}], 1),
            ("Повний фактичний текст відповіді",
             [{"id": "s1", "title": "доказ.pdf", "url": "https://example.test/s1"}], 0),
        )
        for protocol_type in EXPECTED_HASHES:
            for index, (response, documents, absent_count) in enumerate(combinations):
                with self.subTest(protocol_type=protocol_type, combination=index):
                    output = Path(self.temp.name) / f"supplier-result-{protocol_type}-{index}.docx"
                    generator.build_violation_protocol_docx(
                        protocol_type, output, {**BASE_VALUES, "supplier_response": response},
                        "Збережене обґрунтування", [], documents,
                        {"has_written_refusal": False, "has_contract": False,
                         "has_supplier_response": bool(response), "has_supplier_documents": bool(documents),
                         "has_customer_documents": False, "has_civil_code_basis": False,
                         "has_contract_security": False, "has_court_decision": False})
                    text = all_text(output)
                    self.assertIn("Відповідь постачальника на звернення:", text)
                    self.assertIn("Документи подані постачальником:", text)
                    self.assertEqual("Повний фактичний текст відповіді" in text, bool(response))
                    self.assertEqual("доказ.pdf" in text, bool(documents))
                    document = Document(output)
                    rows = [row for table in document.tables for row in table.rows
                            if any(label in " ".join(cell.text for cell in row.cells)
                                   for label in ("Відповідь постачальника", "Документи подані постачальником"))]
                    self.assertEqual(len(rows), 2)
                    row_values = []
                    for row in rows:
                        cells = []
                        seen = set()
                        for cell in row.cells:
                            if cell._tc not in seen:
                                seen.add(cell._tc); cells.append(cell)
                        row_values.append(cells[-1].text.strip())
                        self.assertTrue(all(run.bold for run in cells[0].paragraphs[0].runs if run.text))
                        if cells[-1].text.strip() == "не надано":
                            self.assertTrue(all(run.italic for paragraph in cells[-1].paragraphs
                                                for run in paragraph.runs if run.text))
                    self.assertEqual(sum(value == "не надано" for value in row_values), absent_count)
                    if documents:
                        with zipfile.ZipFile(output) as package:
                            relationships = package.read("word/_rels/document.xml.rels").decode("utf-8")
                        self.assertIn("https://example.test/s1", relationships)

    def test_contract_block_requires_manual_number_and_date(self):
        output = Path(self.temp.name) / "missing_contract.docx"
        with self.assertRaisesRegex(ValueError, "Номер договору.*Дата договору|Дата договору.*Номер договору"):
            generator.build_violation_protocol_docx(
                "decline_p49_1_2", output,
                {**BASE_VALUES, "contract_date": "", "contract_number": ""},
                "Збережене обґрунтування", [], [],
                {"has_contract": True, "has_written_refusal": True,
                 "has_supplier_response": True},
            )
        self.assertFalse(output.exists())

    def test_unknown_placeholder_fails_fast(self):
        custom = Path(self.temp.name) / "unknown.docx"
        shutil.copy2(generator.TEMPLATES["warning"], custom)
        document = Document(custom)
        document.add_paragraph("{{ future_unknown_field }}")
        document.save(custom)
        templates = dict(generator.TEMPLATES); templates["warning"] = custom
        with patch.object(generator, "TEMPLATES", templates):
            with self.assertRaisesRegex(ValueError, "future_unknown_field"):
                self.build("warning")

    def test_unused_unresolved_dative_does_not_block_template(self):
        output = Path(self.temp.name) / "decline_without_dative.docx"
        generator.build_violation_protocol_docx(
            "decline_p49_1_2", output, {**BASE_VALUES, "supplier_name_dative": ""},
            "Збережене обґрунтування", [], [],
            {"has_contract": False, "has_written_refusal": False,
             "has_supplier_response": False, "has_customer_documents": False,
             "has_supplier_documents": False, "has_civil_code_basis": False,
             "has_contract_security": False, "has_court_decision": False},
        )
        self.assertTrue(output.is_file())

    def test_used_unresolved_dative_blocks_warning_template(self):
        output = Path(self.temp.name) / "warning_without_dative.docx"
        with self.assertRaisesRegex(ValueError, "давальний"):
            generator.build_violation_protocol_docx(
                "warning", output, {**BASE_VALUES, "supplier_name_dative": ""},
                "Збережене обґрунтування", [], [],
                {"has_contract": False, "has_written_refusal": False,
                 "has_supplier_response": False, "has_customer_documents": False,
                 "has_supplier_documents": False, "has_civil_code_basis": False,
                 "has_contract_security": False, "has_court_decision": False},
            )
        self.assertFalse(output.exists())

    def test_multiple_documents_keep_stable_order_and_metadata(self):
        output = self.build(customer=[
            {"id": "1", "title": "Перший", "datePublished": "2026-09-01T13:30:45.123+03:00", "url": "https://example.test/1"},
            {"id": "2", "title": "Другий", "datePublished": "2026-09-02T09:05:00Z", "number": "№ 7", "url": "https://example.test/2"},
        ])
        text = all_text(output)
        self.assertLess(text.index("Перший"), text.index("Другий"))
        self.assertNotIn("01.09.2026 13:30", text)
        self.assertNotIn("2026-09-01T", text)
        self.assertNotIn("№ 7", text)
        document = Document(output)
        paragraphs = [p.text for table in document.tables for row in table.rows
                      for cell in row.cells for p in cell.paragraphs]
        self.assertTrue(any(p.startswith("Перший") for p in paragraphs))
        self.assertTrue(any(p.startswith("Другий") for p in paragraphs))
        self.assertFalse(any("Перший" in p and "Другий" in p for p in paragraphs))

    def test_multiline_supplier_response_uses_template_paragraphs_not_double_breaks(self):
        output = Path(self.temp.name) / "multiline-response.docx"
        values = {**BASE_VALUES, "supplier_response": "Перший змістовий абзац.\n\nДругий змістовий абзац."}
        generator.build_violation_protocol_docx(
            "decline_p49_1_2", output, values, "Збережене обґрунтування", [], [],
            {"has_contract": False, "has_written_refusal": False,
             "has_supplier_response": True, "has_customer_documents": False,
             "has_supplier_documents": False, "has_civil_code_basis": False,
             "has_contract_security": False, "has_court_decision": False},
        )
        document = Document(output)
        row = next(row for table in document.tables for row in table.rows
                   if "Відповідь постачальника на звернення" in " ".join(cell.text for cell in row.cells))
        target = row.cells[-1]
        meaningful = [p for p in target.paragraphs if p.text.strip()]
        self.assertEqual([p.text for p in meaningful], [
            "Перший змістовий абзац.", "Другий змістовий абзац."])
        ppr_xml = [p._p.pPr.xml if p._p.pPr is not None else "" for p in meaningful]
        self.assertEqual(ppr_xml[0], ppr_xml[1])

    def test_existing_ukrainian_quote_is_not_doubled(self):
        values = {**BASE_VALUES, "violation_description": "«Опис Замовника.»"}
        output = Path(self.temp.name) / "quoted.docx"
        generator.build_violation_protocol_docx(
            "warning", output, values, "Збережене обґрунтування", [], [],
            {"has_written_refusal": True, "has_contract": True,
             "has_supplier_response": True})
        text = all_text(output)
        self.assertIn("«Опис Замовника.»", text)
        self.assertNotIn("««", text)

    def test_every_nonempty_text_run_has_explicit_times_new_roman(self):
        for protocol_type in EXPECTED_HASHES:
            output = self.build(protocol_type, customer=[{
                "id": "c1", "title": "скарга.pdf",
                "datePublished": "2026-08-31T15:15:58.750734+03:00",
                "url": "https://example.test/customer",
            }], supplier=[{
                "id": "s1", "title": "пояснення.pdf",
                "datePublished": "2026-09-01T10:00:00+03:00",
                "url": "https://example.test/supplier",
            }])
            with zipfile.ZipFile(output) as package:
                from lxml import etree
                xml_parts = [package.read(name) for name in package.namelist()
                             if name.startswith("word/") and name.endswith(".xml")]
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            failures = []
            for xml in xml_parts:
                root = etree.fromstring(xml)
                for run in root.xpath(".//w:r[.//w:t[string-length(.) > 0]]", namespaces=ns):
                    text = "".join(run.xpath(".//w:t/text()", namespaces=ns))
                    fonts = run.find("w:rPr/w:rFonts", ns)
                    values = [] if fonts is None else [fonts.get(qn(f"w:{key}"))
                                                        for key in ("ascii", "hAnsi", "eastAsia", "cs")]
                    if values != ["Times New Roman"] * 4:
                        failures.append((text, values))
            self.assertEqual(failures, [], f"non-Times-New-Roman text runs: {failures[:5]}")

    def test_civil_code_block_is_structured_flag_only_and_keeps_three_links(self):
        enabled = self.build("warning")
        enabled_text = all_text(enabled)
        self.assertEqual(enabled_text.count("роз’яснень Міністерства економіки України"), 1)
        self.assertIn("ч.\u00a05 ст.\u00a0254 Цивільного кодексу України", enabled_text)
        with zipfile.ZipFile(enabled) as package:
            from lxml import etree
            root = etree.fromstring(package.read("word/document.xml"))
            relationships = package.read("word/_rels/document.xml.rels").decode("utf-8")
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        paragraphs = root.xpath(".//w:body/w:p", namespaces=ns)
        texts = ["".join(p.xpath(".//w:t/text()", namespaces=ns)) for p in paragraphs]
        start = next(i for i, text in enumerate(texts)
                     if "роз’яснень Міністерства економіки України" in text)
        heading = next(i for i, text in enumerate(texts)
                       if "За результатами розгляду встановлено" in text)
        normative = paragraphs[start:heading]
        self.assertEqual(len(normative), 8)
        self.assertEqual(sum(len(p.xpath(".//w:hyperlink", namespaces=ns))
                             for p in normative), 3)
        self.assertEqual(sum(bool("".join(p.xpath(".//w:hyperlink//w:t/text()", namespaces=ns)))
                             for p in normative), 3)
        for paragraph in normative:
            self.assertFalse(paragraph.xpath(".//w:highlight | .//w:shd", namespaces=ns))
            for run in paragraph.xpath(".//w:r[.//w:t[string-length(.) > 0]]", namespaces=ns):
                run_properties = run.find(qn("w:rPr"))
                self.assertIsNotNone(run_properties)
                self.assertEqual(run_properties.find(qn("w:sz")).get(qn("w:val")), "24")
                self.assertEqual(run_properties.find(qn("w:szCs")).get(qn("w:val")), "24")
        for url in MINISTRY_URLS:
            self.assertEqual(relationships.count(url.replace("&", "&amp;")), 1)
        self.assertNotIn("{{#if", enabled_text)
        self.assertNotIn("{{/if}}", enabled_text)

        disabled = Path(self.temp.name) / "no-civil-code.docx"
        generator.build_violation_protocol_docx(
            "decline_p49_1_2", disabled, dict(BASE_VALUES),
            "П’ятий день — субота, але цей текст не керує conditional block.", [], [],
            {"has_written_refusal": True, "has_contract": True,
             "has_supplier_response": True, "has_civil_code_basis": False})
        disabled_text = all_text(disabled)
        self.assertNotIn("роз’яснень Міністерства економіки України", disabled_text)
        self.assertNotIn("ч.\u00a05 ст.\u00a0254 Цивільного кодексу України", disabled_text)
        with zipfile.ZipFile(disabled) as package:
            from lxml import etree
            root = etree.fromstring(package.read("word/document.xml"))
        paragraphs = root.xpath(".//w:body/w:p", namespaces=ns)
        texts = ["".join(p.xpath(".//w:t/text()", namespaces=ns)).strip()
                 for p in paragraphs]
        heading = next(i for i, text in enumerate(texts)
                       if "За результатами розгляду встановлено" in text)
        self.assertGreater(heading, 0)
        self.assertTrue(texts[heading - 1])
        heading_properties = paragraphs[heading].find(qn("w:pPr"))
        self.assertTrue(
            heading_properties is None
            or heading_properties.find(qn("w:pageBreakBefore")) is None
        )

    def test_runtime_templates_are_seeded_without_overwriting_operator_version(self):
        runtime = Path(self.temp.name) / "runtime"
        templates = {key: runtime / path.name for key, path in generator.TEMPLATES.items()}
        with patch.object(generator, "TEMPLATE_DIR", runtime), patch.object(generator, "TEMPLATES", templates):
            generator.ensure_runtime_templates()
            self.assertTrue(all(path.exists() for path in templates.values()))
            templates["warning"].write_bytes(b"operator-version")
            generator.ensure_runtime_templates()
            self.assertEqual(templates["warning"].read_bytes(), b"operator-version")

    def test_replacement_rejects_docx_without_required_markers(self):
        runtime = Path(self.temp.name) / "runtime"
        templates = {key: runtime / path.name for key, path in generator.TEMPLATES.items()}
        invalid = Path(self.temp.name) / "invalid.docx"
        Document().save(invalid)
        with patch.object(generator, "TEMPLATE_DIR", runtime), patch.object(generator, "TEMPLATES", templates):
            generator.ensure_runtime_templates()
            original = sha256(templates["warning"])
            with self.assertRaisesRegex(ValueError, "conditional blocks"):
                generator.replace_runtime_template("warning", invalid)
            self.assertEqual(sha256(templates["warning"]), original)

    def test_valid_replacement_keeps_backup_and_reports_active_version(self):
        runtime = Path(self.temp.name) / "runtime"
        templates = {key: runtime / path.name for key, path in generator.TEMPLATES.items()}
        replacement = Path(self.temp.name) / "replacement.docx"
        shutil.copy2(generator.PACKAGED_TEMPLATE_DIR / "warning.docx", replacement)
        with patch.object(generator, "TEMPLATE_DIR", runtime), patch.object(generator, "TEMPLATES", templates):
            generator.ensure_runtime_templates()
            original = sha256(templates["warning"])
            generator.replace_runtime_template("warning", replacement)
            versions = list((runtime / "_versions").glob("warning_*.docx"))
            self.assertEqual(len(versions), 1)
            self.assertEqual(sha256(versions[0]), original)
            item = next(x for x in generator.template_metadata() if x["key"] == "warning")
            self.assertTrue(item["exists"])
            self.assertEqual(item["filename"], "warning.docx")
            self.assertTrue(item["modified_at"])

    def test_replacement_retries_a_transient_windows_lock(self):
        runtime = Path(self.temp.name) / "runtime"
        templates = {key: runtime / path.name for key, path in generator.TEMPLATES.items()}
        replacement = Path(self.temp.name) / "replacement.docx"
        shutil.copy2(generator.PACKAGED_TEMPLATE_DIR / "warning.docx", replacement)
        real_replace = os.replace
        attempts = 0

        def temporarily_locked(source, target):
            nonlocal attempts
            attempts += 1
            if attempts < 3:
                raise PermissionError(5, "Access is denied", str(target))
            return real_replace(source, target)

        with patch.object(generator, "TEMPLATE_DIR", runtime), \
                patch.object(generator, "TEMPLATES", templates), \
                patch.object(generator.os, "replace", side_effect=temporarily_locked), \
                patch.object(generator.time, "sleep"):
            generator.replace_runtime_template("warning", replacement)
        self.assertEqual(attempts, 3)
        self.assertTrue(templates["warning"].is_file())
        self.assertEqual(list(runtime.glob(".warning.*.tmp.docx")), [])

    def test_replacement_lock_keeps_active_template_and_cleans_partial_files(self):
        runtime = Path(self.temp.name) / "runtime"
        templates = {key: runtime / path.name for key, path in generator.TEMPLATES.items()}
        replacement = Path(self.temp.name) / "replacement.docx"
        shutil.copy2(generator.PACKAGED_TEMPLATE_DIR / "warning.docx", replacement)
        with patch.object(generator, "TEMPLATE_DIR", runtime), patch.object(generator, "TEMPLATES", templates):
            generator.ensure_runtime_templates()
            original = sha256(templates["warning"])
            with patch.object(generator.os, "replace", side_effect=PermissionError(5, "Access is denied")), \
                    patch.object(generator.time, "sleep"):
                with self.assertRaisesRegex(PermissionError, "заблокований Windows"):
                    generator.replace_runtime_template("warning", replacement)
            self.assertEqual(sha256(templates["warning"]), original)
            self.assertEqual(list(runtime.glob(".warning.*.tmp.docx")), [])
            self.assertEqual(list((runtime / "_versions").glob("warning_*.docx")), [])

    def test_regeneration_uses_atomic_sibling_and_reports_external_lock(self):
        output = self.build()
        original = output.read_bytes()
        real_replace = os.replace

        def locked_replace(source, target):
            if Path(target) == output:
                raise PermissionError(13, "file is used by another process", str(target))
            return real_replace(source, target)

        with patch.object(generator.os, "replace", side_effect=locked_replace):
            with self.assertRaisesRegex(PermissionError, "Закрийте його у Microsoft Word"):
                generator.build_violation_protocol_docx(
                    "warning", output,
                    {**BASE_VALUES, "protocol_number": "TEST-2"},
                    "ОСТАТОЧНЕ ОБҐРУНТУВАННЯ УО ДОСЛІВНО", [], [],
                    {"has_written_refusal": True, "has_contract": True, "has_contract_security": True,
                     "has_civil_code_basis": True, "has_court_decision": True,
                     "has_supplier_response": True},
                )
        self.assertEqual(output.read_bytes(), original)
        self.assertEqual(list(output.parent.glob(f".{output.stem}.*.tmp.docx")), [])


if __name__ == "__main__":
    unittest.main()
