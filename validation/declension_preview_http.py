"""Synthetic HTTP + shared declension + new-DOCX regression; no live storage."""
import base64,gc,hashlib,http.client,json,sys,threading,unittest,urllib.parse,os
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parent))
os.environ['PQM_RELEASE_SCHEMA_ONLY']='1'
import web_smoke as web
import declension,declension_overrides,violation_protocol_docx as docx_generator
from docx import Document
s=web.server

class PreviewRegression(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  web.fixture();cls.path=Path(web.TEMP.name)/'preview.csv';cls.store=declension.OverrideStore(cls.path)
  cls.patchers=[patch.object(declension,'DEFAULT_OVERRIDE_STORE',cls.store),patch.object(s,'list_overrides',lambda:declension_overrides.list_overrides(cls.path)),patch.object(s,'save_override',lambda payload,record_id=None:declension_overrides.save_override(payload,record_id,cls.path))]
  for p in cls.patchers:p.start()
  cls.http=s.ThreadingHTTPServer(('127.0.0.1',0),s.Handler);cls.thread=threading.Thread(target=cls.http.serve_forever,daemon=True);cls.thread.start();cls.port=cls.http.server_address[1]
 @classmethod
 def tearDownClass(cls):
  cls.http.shutdown();cls.http.server_close();cls.thread.join(5)
  for p in reversed(cls.patchers):p.stop()
  gc.collect();web.TEMP.cleanup()
 def setUp(self):
  self.path.write_text('entity_type;original;genitive;dative;accusative;comment\n',encoding='utf-8');self.store._signature=None
 def request(self,path,role='admin',method='GET',payload=None):
  headers={}
  if role:headers['Authorization']='Basic '+base64.b64encode(('fixture-'+role+':'+web.PASSWORD).encode()).decode()
  if payload is not None:headers['Content-Type']='application/json'
  con=http.client.HTTPConnection('127.0.0.1',self.port,timeout=10);con.request(method,path,body=json.dumps(payload) if payload is not None else None,headers=headers);response=con.getresponse();status=response.status;data=json.loads(response.read());con.close();return status,data
 def preview(self,name,entity_type='legal_entity',role='admin'):
  return self.request('/api/declension-overrides?'+urllib.parse.urlencode({'preview':'1','original':name,'entity_type':entity_type}),role)
 def test_preview_equals_generator_resolver_and_never_writes(self):
  before=self.path.read_bytes();name='КОМУНАЛЬНЕ ПІДПРИЄМСТВО «СИНТЕТИЧНЕ»';status,data=self.preview(name);self.assertEqual(status,200)
  for case in ('genitive','dative','accusative'):self.assertEqual(data['preview'][case],declension.decline_name(name,'legal_entity',case,self.store).as_dict())
  self.assertEqual(data['preview']['genitive']['source'],'automatic');self.assertEqual(data['preview']['dative']['value'],'');self.assertEqual(self.path.read_bytes(),before)
 def test_prat_and_dative_remain_unresolved_without_invented_algorithm(self):
  _,data=self.preview('ПРАТ «СИНТЕТИЧНА ЮРІЯ»');self.assertTrue(all(v['source']=='unresolved' and not v['value'] for v in data['preview'].values()))
 def test_verified_override_priority_and_quote_identity(self):
  row=declension_overrides.save_override({'entity_type':'legal_entity','original':'ТОВ "СИНТЕТИЧНЕ"','genitive':'ПЕРЕВІРЕНИЙ РОДОВИЙ','dative':'ПЕРЕВІРЕНИЙ ДАВАЛЬНИЙ','accusative':'','comment':'ручна перевірка'},path=self.path)
  _,data=self.preview('ТОВ «СИНТЕТИЧНЕ»');self.assertEqual(data['item']['id'],row['id']);self.assertEqual(data['preview']['dative']['source'],'override');self.assertEqual(data['preview']['genitive']['value'],'ПЕРЕВІРЕНИЙ РОДОВИЙ')
 def test_roles_keep_existing_read_and_write_contract(self):
  for role in ('admin','officer','viewer',None):
   baseline,_=self.request('/api/declension-overrides',role);status,_=self.preview('ТОВ «СИНТЕТИЧНЕ»',role=role);self.assertEqual(status,baseline)
  status,_=self.request('/api/admin/declension-overrides','viewer','POST',{'entity_type':'legal_entity','original':'ТОВ «СИНТЕТИЧНЕ»'});self.assertEqual(status,403)
 def test_validation_and_default_list_compatibility(self):
  self.assertEqual(self.preview('',role='admin')[0],400);self.assertEqual(self.preview('X','invalid')[0],400)
  status,data=self.request('/api/declension-overrides');self.assertEqual(status,200);self.assertEqual(data,{'items':[]})
 def test_declension_status_readback_does_not_refresh_prozorro(self):
  with patch.object(s,'violation_report_detail',return_value={'id':'synthetic','protocol_readiness':{}}) as detail:
   status,_=self.request('/api/violation-reports/synthetic?refresh=0');self.assertEqual(status,200);detail.assert_called_once_with('synthetic',refresh=False)
  with patch.object(s,'violation_report_detail',return_value={'id':'synthetic'}) as detail:
   status,_=self.request('/api/violation-reports/synthetic');self.assertEqual(status,200);detail.assert_called_once_with('synthetic')
 def test_blank_required_dative_never_resolves_after_save_or_delete(self):
  name='ТОВ «СИНТЕТИЧНИЙ ОБОВЯЗКОВИЙ»';payload={'entity_type':'legal_entity','original':name,'genitive':'','dative':'ПЕРЕВІРЕНА ФОРМА','accusative':'','comment':'required_case=dative'}
  status,data=self.request('/api/admin/declension-overrides','admin','POST',payload);self.assertEqual(status,201);identifier=data['item']['id']
  for blank in ('','   ','\t'):
   status,_=self.request('/api/admin/declension-overrides/'+identifier,'admin','PATCH',{**payload,'dative':blank});self.assertEqual(status,200)
   _,data=self.preview(name);self.assertEqual(data['preview']['dative']['status'],'unresolved');self.assertEqual(data['preview']['dative']['source'],'unresolved');self.assertEqual(data['preview']['dative']['value'],'');self.assertTrue(data['item']['needs_completion'])
 def test_unconfirmed_auto_stays_auto_until_explicit_save(self):
  name='КОМУНАЛЬНЕ ПІДПРИЄМСТВО «СИНТЕТИЧНИЙ ПРІОРИТЕТ»';payload={'entity_type':'legal_entity','original':name,'genitive':'','dative':'ПЕРЕВІРЕНИЙ ДАВАЛЬНИЙ','accusative':'','comment':'перевірено'}
  status,row=self.request('/api/admin/declension-overrides','admin','POST',payload);self.assertEqual(status,201);identifier=row['item']['id']
  before=self.path.read_bytes();_,data=self.preview(name);self.assertEqual(self.path.read_bytes(),before);self.assertEqual(data['preview']['genitive']['source'],'automatic');self.assertEqual(data['item']['genitive'],'')
  value=data['preview']['genitive']['value'];self.request('/api/admin/declension-overrides/'+identifier,'admin','PATCH',{**payload,'genitive':value})
  _,data=self.preview(name);self.assertEqual(data['preview']['genitive']['source'],'override');self.assertEqual(data['preview']['genitive']['value'],value)
 def test_saved_form_reopens_and_reaches_only_new_docx(self):
  name='ПРАТ «СИНТЕТИЧНЕ»';old=Path(web.TEMP.name)/'signed-historical.docx';old.write_bytes(b'synthetic immutable signed history');before=hashlib.sha256(old.read_bytes()).hexdigest()
  status,data=self.request('/api/admin/declension-overrides','officer','POST',{'entity_type':'legal_entity','original':name,'genitive':'','dative':'ПРАТ «ПЕРЕВІРЕНО»','accusative':'','comment':'перевірено'});self.assertEqual(status,201)
  _,preview=self.preview(name);self.assertEqual(preview['preview']['dative']['value'],'ПРАТ «ПЕРЕВІРЕНО»');self.assertEqual(preview['preview']['dative']['source'],'override')
  template=Path(web.TEMP.name)/'template.docx';doc=Document();doc.add_paragraph('{{supplier_name_dative}}');table=doc.add_table(rows=1, cols=2);table.cell(0,0).text='Обґрунтування рішення';table.cell(0,1).text='{{decision_justification}}';doc.save(template)
  output=Path(web.TEMP.name)/'new.docx'
  with patch.object(docx_generator,'ensure_runtime_templates'),patch.dict(docx_generator.TEMPLATES,warning=template):
   docx_generator.build_violation_protocol_docx('warning',output,{'supplier_name_dative':s.decline_name(name,'legal_entity','dative').value},'Синтетичне обґрунтування')
  self.assertIn('ПРАТ «ПЕРЕВІРЕНО»','\n'.join(p.text for p in Document(output).paragraphs));self.assertEqual(hashlib.sha256(old.read_bytes()).hexdigest(),before)

if __name__=='__main__':unittest.main(verbosity=2)
