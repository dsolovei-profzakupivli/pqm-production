"""Baseline-compatible admission rules; synthetic, read-only fixture execution."""
import builtins
import io
import json
import sqlite3
import unittest
import warnings
from unittest.mock import patch
from zipfile import ZipFile
import admission_resolver as resolver
import supplier_selection
import supplier_xlsx


class AdmissionResolverTests(unittest.TestCase):
    def setUp(self):
        self.con=sqlite3.connect(':memory:');self.con.row_factory=sqlite3.Row
        self.addCleanup(self.con.close)
        self.con.executescript('''CREATE TABLE submissions(id TEXT,supplier_code TEXT,qualification_id TEXT);
          CREATE TABLE qualifications(id TEXT,submission_id TEXT,status TEXT,raw_json TEXT);
          CREATE TABLE supplier_edr_verification_events(supplier_code TEXT,event_type TEXT,
          occurred_at TEXT,source_submission_id TEXT,snapshot_json TEXT);''')
        self.stamp='2026-10-14T14:15:21.004492+03:00'
        self.raw=json.dumps({'id':'q','submissionID':'s','status':'active','date':self.stamp})
        self.con.execute('INSERT INTO submissions VALUES(?,?,?)',('s','001','q'))
        self.con.execute('INSERT INTO qualifications VALUES(?,?,?,?)',('q','s','active',self.raw))

    def test_admission_without_local_protocol_fields(self):
        self.assertEqual(resolver.current_admissions(self.con,{'001'})[0]['occurred_at'],self.stamp)

    def test_identity_status_and_timestamp_are_all_required(self):
        for field,value in [('id','other'),('submissionID','other'),('status','pending'),('date','invalid'),('date','2026-10-14T14:15:21')]:
            raw=json.loads(self.raw);raw[field]=value
            self.assertEqual(resolver.qualification_stamp('active','q','s',json.dumps(raw)),'')
        for state in ['pending','unsuccessful','cancelled']:
            self.assertEqual(resolver.qualification_stamp(state,'q','s',self.raw),'')

    def test_full_precision_offsets_and_rejected_invalid_offsets(self):
        self.assertEqual(resolver.source_instant(self.stamp).microsecond,4492)
        self.assertEqual(resolver.source_instant('2026-10-14T14:15:21.999999+03:00').microsecond,999999)
        self.assertEqual(resolver.source_instant(self.stamp),resolver.source_instant('2026-10-14T11:15:21.004492Z'))
        for stamp in ['2026-10-14T14:15:21+15:00','2026-10-14T14:15:21+14:01','2026-02-30T14:15:21Z']:
            self.assertIsNone(resolver.source_instant(stamp))

    def test_historical_admission_survives_status_change_with_proof_only(self):
        self.con.execute("UPDATE qualifications SET status='cancelled'")
        for proof,kind,sid in [({'submission_id':'s','decision':'admit'},'admission','s'),({},'admission','s'),({'submission_id':'s','decision':'admit'},'exclusion','s'),({'submission_id':'other','decision':'admit'},'admission','s')]:
            self.con.execute('DELETE FROM supplier_edr_verification_events')
            self.con.execute('INSERT INTO supplier_edr_verification_events VALUES(?,?,?,?,?)',('001',kind,self.stamp,sid,json.dumps(proof)))
            events=resolver.admission_events(self.con,{'001'})
            self.assertEqual(len(events),int(kind=='admission' and proof.get('submission_id')=='s' and proof.get('decision')=='admit'))

    def test_read_only_no_unfinished_edr_or_transport_imports(self):
        self.con.commit();self.con.execute('PRAGMA query_only=ON');before=self.con.total_changes
        original=builtins.__import__
        def guarded(name,*args,**kwargs):
            if name in {'edr_admission_read_model','edr_sync_v2','google_writer_execution','sandbox_runtime'}:
                raise AssertionError('Unexpected runtime dependency: '+name)
            return original(name,*args,**kwargs)
        with patch('builtins.__import__',side_effect=guarded):
            self.assertEqual(len(resolver.admission_events(self.con,{'001'})),1)
        self.assertEqual(self.con.total_changes,before)

    def test_xlsx_literal_crlf_cr_lf_unicode_and_formula_safety(self):
        from openpyxl import load_workbook
        literals=['ФОП "ПРОДАН НАТАЛІЯ ВІКТОРІВНА\r\n"','=HYPERLINK("bad")\r\nтекст','@unsafe\rvalue','Україна\nтекст','&#13; literal','_x000D_ literal']
        for literal in literals:
            payload=supplier_xlsx.workbook([{'values':['001',literal,self.stamp,'Активний','mail@example.invalid','+38000123']}])
            with warnings.catch_warnings():
                warnings.simplefilter('ignore');ws=load_workbook(io.BytesIO(payload),read_only=True).active
            self.assertEqual(ws['B2'].value,literal);self.assertEqual(ws['B2'].data_type,'s')
            self.assertEqual(ws['A2'].value,'001');self.assertEqual(ws['F2'].value,'+38000123')
            with ZipFile(io.BytesIO(payload)) as z:
                xml=z.read('xl/worksheets/sheet1.xml');self.assertNotIn(b'<f>',xml)
                if '\r' in literal:self.assertIn(b'&#13;',xml)


if __name__=='__main__':unittest.main()
